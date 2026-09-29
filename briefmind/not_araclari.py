"""
not_araclari.py — notun ve kayitlarin uzerinde calisan yardimcilar (arayuzden bagimsiz):
toplantilar arasi arama, transkriptte bir ana gitme, nottan aksiyon ayiklama, kisiye ozel e-posta,
Word'e aktarma, katilimci bilgilendirme metni.
"""
import json
import os
import re

import motor as motor_mod
from sozluk import normalize

BILGILENDIRME = ("Bilgi: Bu toplantıda BriefMind ile not alınmaktadır. Ses kaydedilmez; konuşmalar kurum içi "
                 "sistemlerde metne çevrilip özetlenir ve not katılımcılarla paylaşılır. Not alınmasını "
                 "istemeyen katılımcı lütfen belirtsin.")


def _sn(ts):
    return motor_mod._sn(ts)


def _katla(metin):
    """Arama icin: buyuk/kucuk harf ve Turkce karakter farkini yok say ('ILETIN' ~ 'iletin', 'sube' ~ 'şube')."""
    m = normalize(metin or "")
    for a, b in zip("çğıöşüâî", "cgiosuai"):
        m = m.replace(a, b)
    return m


def _meta(klasor):
    try:
        with open(os.path.join(klasor, "meta.json"), encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def transkript_satirlari(klasor):
    """Toplantinin satirlari: once parca dosyalari (duzeltilmis metin), yoksa altyazi.jsonl."""
    satirlar = []
    pk = os.path.join(klasor, "parcalar")
    if os.path.isdir(pk):
        for ad in sorted(os.listdir(pk)):
            try:
                with open(os.path.join(pk, ad), encoding="utf-8") as f:
                    satirlar += json.load(f).get("satirlar") or []
            except Exception:
                continue
    if not satirlar and os.path.exists(os.path.join(klasor, "altyazi.jsonl")):
        with open(os.path.join(klasor, "altyazi.jsonl"), encoding="utf-8") as f:
            for ham in f:
                try:
                    satirlar.append(json.loads(ham))
                except Exception:
                    continue
    return satirlar


def ana_git(klasor, ts, pencere_sn=90):
    """ts cevresindeki (±pencere) satirlar: [(ts, konusmaci, metin, hedef_mi)]."""
    hedef = _sn(ts)
    satirlar = [s for s in transkript_satirlari(klasor) if abs(_sn(s.get("ts")) - hedef) <= pencere_sn]
    if not satirlar:
        return []
    en_yakin = min(satirlar, key=lambda s: abs(_sn(s.get("ts")) - hedef))
    return [(s.get("ts"), s.get("speaker", "?"), s.get("text", ""), s is en_yakin) for s in satirlar]


def ara(sorgu, kok=None, en_fazla=200):
    """Tum toplantilarin transkript ve notlarinda arar (buyuk/kucuk harf ve Turkce karakter duyarsiz).
    Sorgudaki TUM kelimeler ayni satirda gecmeli. Dondurur: [{klasor, tarih, baslik, tur, ts, kim, metin}]."""
    kok = kok or motor_mod.KOK
    kelimeler = _katla(sorgu).split()
    if not kelimeler or not os.path.isdir(kok):
        return []
    sonuc = []
    for ad in sorted(os.listdir(kok), reverse=True):
        k = os.path.join(kok, ad)
        meta = _meta(k)
        if not meta:
            continue
        ortak = {"klasor": k, "tarih": meta.get("tarih", ""), "baslik": meta.get("baslik", ad)}
        not_yol = os.path.join(k, "not.md")
        if os.path.exists(not_yol):
            with open(not_yol, encoding="utf-8") as f:
                for satir in f:
                    if all(w in _katla(satir) for w in kelimeler):
                        ts = re.search(r"⏱(\d{1,2}:\d{2}:\d{2})", satir)
                        sonuc.append({**ortak, "tur": "not", "ts": ts.group(1) if ts else "", "kim": "",
                                      "metin": satir.strip().strip("|-# ")})
        for s in transkript_satirlari(k):
            if all(w in _katla(s.get("text", "")) for w in kelimeler):
                sonuc.append({**ortak, "tur": "transkript", "ts": s.get("ts", ""), "kim": s.get("speaker", "?"),
                              "metin": s.get("text", "")})
        if len(sonuc) >= en_fazla:
            break
    return sonuc[:en_fazla]


def aksiyonlari_ayikla(not_md):
    """Nottaki '## Aksiyonlar' tablosu -> [{madde, sorumlu, tarih, kaynak}] (duzenlenmis notta da calisir)."""
    aksiyonlar, bolum = [], ""
    for satir in not_md.splitlines():
        s = satir.strip()
        if s.startswith("#"):
            bolum = s.lower()
            continue
        if "aksiyon" not in bolum or not s.startswith("|"):
            continue
        h = [x.strip() for x in s.strip("|").split("|")]
        if len(h) < 4 or h[0] in ("#", "") or set(h[0]) <= set("-: "):
            continue
        aksiyonlar.append({"madde": h[1], "sorumlu": h[2], "tarih": h[3], "kaynak": h[4] if len(h) > 4 else ""})
    return aksiyonlar


def kisiye_ozel_epostalar(not_md, baslik, tarih):
    """Sorumlusu belli her kisi icin (ad, konu, govde): yalnizca o kisinin isleri + notun ozeti."""
    import llm
    aks = aksiyonlari_ayikla(not_md)
    ozet = ""
    m = re.search(r"## Özet\s*\n(.*?)(\n## |\Z)", not_md, flags=re.S)
    if m:
        ozet = m.group(1).strip()
    epostalar = []
    for ad, isler in llm.kisi_bazli(aks):
        if ad in ("belirsiz", "-", "?"):
            continue
        tarihler = {a["madde"]: a["tarih"] for a in aks}
        govde = [f"Merhaba {ad.split()[0]},", "", f"{tarih} tarihli \"{baslik}\" toplantısında sana düşen işler:", ""]
        govde += [f"- {x}" + (f" (tarih: {tarihler[x]})" if tarihler.get(x) not in (None, "", "-") else "")
                  for x in isler]
        govde += ["", "Toplantı özeti:", ozet or "-", "", "Tam not ekte ya da ortak klasörde."]
        epostalar.append((ad, f"Toplantı aksiyonların: {baslik} ({tarih})", "\n".join(govde)))
    return epostalar


def word_kaydet(not_md, yol):
    """Notu .docx olarak kaydeder (python-docx). Kurulu degilse ImportError."""
    import docx
    belge = docx.Document()
    tablo, sutun = None, 0
    for ham in not_md.splitlines():
        s = ham.strip()
        if s.startswith("|"):
            h = [x.strip() for x in s.strip("|").split("|")]
            if all(set(x) <= set("-: ") for x in h):
                continue
            if tablo is None:
                sutun = len(h)
                tablo = belge.add_table(rows=0, cols=sutun)
                tablo.style = "Table Grid"
            hucreler = tablo.add_row().cells
            for i, x in enumerate(h[:sutun]):
                hucreler[i].text = x
            continue
        tablo = None
        if s.startswith("# "):
            belge.add_heading(s[2:], level=0)
        elif s.startswith("## "):
            belge.add_heading(s[3:], level=1)
        elif s.startswith(("- ", "* ")):
            belge.add_paragraph(s[2:], style="List Bullet")
        elif s:
            belge.add_paragraph(s)
    belge.save(yol)
    return yol

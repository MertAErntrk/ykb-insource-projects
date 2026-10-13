"""Kaynaga gitme, kisi bazli liste, sablon, arama, e-posta, Word, saklama suresi, toplantiya soru."""
import datetime as dt
import json
import os

import llm
import motor as motor_mod
import not_araclari

SATIRLAR = [
    {"ts": "10:02:00", "speaker": "Selin Kara", "text": "MVP prototipini ben hazırlarım, edite edilebilir olsun"},
    {"ts": "10:05:30", "speaker": "Emre Sancak", "text": "Ara tabloyu Excel olarak bana iletin, Deniz Can ile incelerim"},
    {"ts": "10:09:10", "speaker": "Kemal Tuncer", "text": "Follow-up toplantılarını 2-3 haftada bir yapalım"},
]
LISTELER = {"kararlar": ["Follow-up toplantıları 2-3 haftada bir yapılacak"],
            "aksiyonlar": [{"madde": "Edite edilebilir MVP prototipi hazırlamak", "sorumlu": "Selin Kara", "tarih": "-"},
                           {"madde": "Ara tabloyu Excel olarak incelemek", "sorumlu": "Emre Sancak, Deniz Can Öztürk",
                            "tarih": "-"},
                           {"madde": "Örnekleri hazırlamak", "sorumlu": "Selin Kara", "tarih": "-"},
                           {"madde": "Kayıtları iletmek", "sorumlu": "belirsiz", "tarih": "-"}],
            "acik_sorular": []}


def _klasor(tmp_path, durum="tamam", tarih="2026-06-01"):
    k = tmp_path / "toplantilar" / f"{tarih}_strateji"
    (k / "parcalar").mkdir(parents=True)
    (k / "meta.json").write_text(json.dumps({"baslik": "Strateji", "tarih": tarih, "durum": durum, "parca": 1}),
                                 encoding="utf-8")
    (k / "parcalar" / "parca_001.json").write_text(json.dumps(
        {"sira": 1, "baslangic": "10:02:00", "bitis": "10:09:10", "satirlar": SATIRLAR, "duzeltmeler": [],
         "ozet": {"ozet": "MVP ve follow-up konuşuldu."}}, ensure_ascii=False), encoding="utf-8")
    (k / "altyazi.jsonl").write_text("\n".join(json.dumps(s, ensure_ascii=False) for s in SATIRLAR), encoding="utf-8")
    return k


def test_kaynak_bulucu_ve_not_markdown():
    bul = motor_mod.kaynak_bulucu(SATIRLAR)
    assert bul("Edite edilebilir MVP prototipi hazırlamak", "Selin Kara") == "10:02:00"
    assert bul("Follow-up toplantıları 2-3 haftada bir yapılacak") == "10:09:10"
    assert bul("Bütçe onayı alınacak") is None
    md = llm.not_markdown({"ozet": "Özet paragrafı burada.", "sonraki_adim": "-"}, LISTELER, bul)
    assert "- Follow-up toplantıları 2-3 haftada bir yapılacak ⏱10:09:10" in md
    assert "| # | Madde | Sorumlu | Tarih | Kaynak |" in md and "| ⏱10:02:00 |" in md
    assert "## Kişiye göre iş listesi" in md
    assert "- Selin Kara (2): Edite edilebilir MVP prototipi hazırlamak; Örnekleri hazırlamak" in md
    assert "- Deniz Can Öztürk (1): Ara tabloyu Excel olarak incelemek" in md
    assert md.index("- belirsiz (1)") > md.index("- Selin Kara")          # belirsiz en sonda
    # dayanak kontrolu kisi listesini ve kaynak isaretlerini bozmaz
    yeni, atilan = motor_mod.dayanak_kontrolu("# T\n\n" + md, " ".join(s["text"] for s in SATIRLAR))
    assert "## Kişiye göre iş listesi" in yeni and "⏱10:09:10" in yeni


def test_sablon_talimati_baglama_eklenir(monkeypatch):
    goren = []

    def sahte(mesajlar, max_tokens, sema=None, effort="medium", temperature=0.6, dusunme=True):
        goren.append(mesajlar[1]["content"])
        return json.dumps({"ozet": "Toplantıda ilerleme ve engeller ayrıntılı olarak konuşuldu, plan yapıldı.",
                           "sonraki_adim": "-", "kararlar": [], "aksiyonlar": [], "acik_sorular": []}), "stop"

    monkeypatch.setattr(llm, "sor", sahte)
    monkeypatch.setattr(llm, "token_say", llm.token_tahmin)
    llm.birlestir([{"ozet": "a", "konular": [], "kararlar": [], "aksiyonlar": [], "acik_sorular": []}],
                  "Toplantı: x", sablon="haftalik")
    assert any("ilerleme, engeller" in g for g in goren)


def test_arama_ana_git_aksiyon_ayikla_eposta_word(tmp_path):
    k = _klasor(tmp_path)
    sonuc = not_araclari.ara("excel ILETIN", kok=str(tmp_path / "toplantilar"))
    assert len(sonuc) == 1 and sonuc[0]["kim"] == "Emre Sancak" and sonuc[0]["ts"] == "10:05:30"
    pencere = not_araclari.ana_git(str(k), "10:05:30", pencere_sn=120)
    assert [x[0] for x in pencere] == ["10:05:30"] or any(x[3] and x[0] == "10:05:30" for x in pencere)
    md = llm.not_markdown({"ozet": "Özet.", "sonraki_adim": "-"}, LISTELER, motor_mod.kaynak_bulucu(SATIRLAR))
    aks = not_araclari.aksiyonlari_ayikla(md)
    assert len(aks) == 4 and aks[0]["sorumlu"] == "Selin Kara" and aks[0]["kaynak"] == "⏱10:02:00"
    epostalar = not_araclari.kisiye_ozel_epostalar(md, "Strateji", "2026-09-23")
    adlar = [e[0] for e in epostalar]
    assert adlar[0] == "Selin Kara" and "belirsiz" not in adlar and "Deniz Can Öztürk" in adlar
    assert "Örnekleri hazırlamak" in epostalar[0][2] and "Kayıtları iletmek" not in epostalar[0][2]
    yol = not_araclari.word_kaydet("# T\n\n" + md, str(tmp_path / "n.docx"))
    import docx
    metin = "\n".join(p.text for p in docx.Document(yol).paragraphs)
    assert "Kişiye göre iş listesi" in metin and len(docx.Document(yol).tables) == 1


def test_saklama_suresi(tmp_path):
    kok = str(tmp_path / "toplantilar")
    eski = _klasor(tmp_path, tarih="2026-06-01")
    yeni = _klasor(tmp_path, tarih="2026-09-25")
    (eski / "not.md").write_text("# not", encoding="utf-8")
    bugun = dt.date(2026, 9, 29)
    assert motor_mod.saklama_uygula(0, kok, bugun) == 0                    # kapali
    assert motor_mod.saklama_uygula(90, kok, bugun) == 1
    assert not (eski / "altyazi.jsonl").exists() and (yeni / "altyazi.jsonl").exists()
    p = json.loads((eski / "parcalar" / "parca_001.json").read_text(encoding="utf-8"))
    assert p["satirlar"] == [] and p["ozet"]                               # ozet kalir
    assert (eski / "not.md").exists()
    assert json.loads((eski / "meta.json").read_text(encoding="utf-8"))["transkript_silindi"] == "2026-09-29"
    assert motor_mod.saklama_uygula(90, kok, bugun) == 0                   # ikinci kez dokunmaz
    # yarim (durumu 'tamam' olmayan) toplanti silinmez
    yarim = _klasor(tmp_path, durum="devam", tarih="2026-01-01")
    assert motor_mod.saklama_uygula(90, kok, bugun) == 0 and (yarim / "altyazi.jsonl").exists()


def test_kaydi_acmak_durumu_bozmaz(tmp_path):
    k = _klasor(tmp_path)
    m = motor_mod.Motor.yukle(str(k))
    m.kapat()
    assert json.loads((k / "meta.json").read_text(encoding="utf-8"))["durum"] == "tamam"


def test_toplantiya_soru_ilgili_bolumu_secer(monkeypatch):
    goren = {}

    def sahte(mesajlar, max_tokens, sema=None, effort="medium", temperature=0.6, dusunme=True):
        goren["icerik"] = mesajlar[1]["content"]
        return "Selin Kara hazırlayacak ⏱10:02:00", "stop"

    monkeypatch.setattr(llm, "sor", sahte)
    monkeypatch.setattr(llm, "token_say", llm.token_tahmin)
    monkeypatch.setattr(llm, "BAGLAM_PENCERESI", 2600)                    # tek bolum sigacak kadar dar
    bolumler = [("09:00–09:10", "[09:01:00] Ali: Bütçe konuşuldu. " * 40),
                ("10:00–10:10", "[10:02:00] Selin Kara: MVP prototipini ben hazırlarım"),
                ("11:00–11:10", "[11:01:00] Ayşe: Kampanya takvimi. " * 40)]
    cevap = llm.toplantiya_sor("MVP prototipini kim hazırlayacak?", bolumler, "Toplantı: x")
    assert "Selin Kara" in cevap and "MVP prototipini" in goren["icerik"]
    assert "Kampanya takvimi" not in goren["icerik"]


# ---------- C3: konu etiketleri (KC3: Konu sutunu HER ZAMAN sonda, '[konu: X]' damgasi) ----------
NOT_ALTI_SUTUN = """# Haftalık — 2026-10-12

## Özet
Özet.

## Kararlar
- Bütçe onaylandı [konu: Veri Ambarı] ⏱10:02:00
- Takvim kaydı [konu: Yapay Zekâ] ⏱10:09:10

## Aksiyonlar
| # | Madde | Sorumlu | Tarih | Kaynak | Konu |
|---|---|---|---|---|---|
| 1 | Şemayı hazırlamak | Selin Kara | 2026-10-20 | ⏱10:02:00 | Veri Ambarı |
| 2 | Model karşılaştırmak | Emre Sancak (?) | - | ⏱10:05:30 | Yapay Zekâ |
| 3 | Yedek planı yazmak | belirsiz | - | - | - |
"""
NOT_ESKI_BES_SUTUN = """## Aksiyonlar
| # | Madde | Sorumlu | Tarih | Kaynak |
|---|---|---|---|---|
| 1 | Şemayı hazırlamak | Selin Kara | 2026-10-20 | ⏱10:02:00 |
"""


def test_aksiyon_ayikla_konu_sutunu_adla():
    aks = not_araclari.aksiyonlari_ayikla(NOT_ALTI_SUTUN)
    assert [a["konu"] for a in aks] == ["Veri Ambarı", "Yapay Zekâ", None]       # '-' -> None
    assert aks[0]["madde"] == "Şemayı hazırlamak" and aks[0]["kaynak"] == "⏱10:02:00"
    assert aks[1]["sahip"]["durum"] == "belirsiz" and aks[2]["sahip"]["durum"] == "yok"


def test_aksiyon_ayikla_eski_bes_sutun_ve_basliksiz():
    aks = not_araclari.aksiyonlari_ayikla(NOT_ESKI_BES_SUTUN)
    assert len(aks) == 1 and aks[0]["konu"] is None and aks[0]["kaynak"] == "⏱10:02:00"
    basliksiz = "## Aksiyonlar\n| 1 | İş A | Selin Kara | 2026-10-20 | ⏱10:02:00 | Veri Ambarı |\n| 2 | İş B | Emre | - |\n"
    aks = not_araclari.aksiyonlari_ayikla(basliksiz)
    assert [a["madde"] for a in aks] == ["İş A", "İş B"] and aks[0]["konu"] == "Veri Ambarı" and aks[1]["konu"] is None
    assert aks[1]["kaynak"] == ""


def test_aksiyon_ayikla_sutun_sirasi_serbest_ve_damga_siyrilir():
    md = ("## Aksiyonlar\n| # | Konu | Madde | Sorumlu | Tarih |\n|---|---|---|---|---|\n"
          "| 1 | Bütçe | Teklif almak [konu: Satın alma] | Selin Kara | - |\n"
          "| 2 | - | Rapor yazmak [konu: Satın alma] | Selin Kara | - |\n")
    aks = not_araclari.aksiyonlari_ayikla(md)
    assert aks[0]["madde"] == "Teklif almak" and aks[0]["konu"] == "Bütçe"        # Konu hucresi damgadan once gelir
    assert aks[1]["madde"] == "Rapor yazmak" and aks[1]["konu"] == "Satın alma"   # hucre bos -> damga


def test_durum_sablonu_tek_duz_aksiyon_tablosu():
    md = """# Durum — 2026-10-12

## Konular
### Veri Ambarı
Durum: Şema tamam. Risk: Takvim sıkışık.
#### Kararlar
- Şema onaylandı ⏱10:02:00
### Genel
#### Açık sorular
- Bütçe kimde?

## Aksiyonlar
| # | Madde | Sorumlu | Tarih | Kaynak | Konu |
|---|---|---|---|---|---|
| 1 | Şemayı yayınlamak | Selin Kara | - | ⏱10:02:00 | Veri Ambarı |
| 2 | Bütçeyi sormak | Emre Sancak | - | - | - |
"""
    aks = not_araclari.aksiyonlari_ayikla(md)
    assert [a["madde"] for a in aks] == ["Şemayı yayınlamak", "Bütçeyi sormak"]
    assert [a["konu"] for a in aks] == ["Veri Ambarı", None]
    assert not_araclari.aksiyon_sirasini_yenile(md).count("| 1 |") == 1


def test_kisiye_gore_yenile_konuyu_gosterir_ve_adi_korur():
    md = ("## Aksiyonlar\n| # | Madde | Sorumlu | Tarih | Kaynak | Konu |\n|---|---|---|---|---|---|\n"
          "| 1 | A işi | Selin Kara | - | - | Veri Ambarı |\n| 2 | B işi | Selin Kara | - | - | - |\n"
          "| 3 | C işi | Emre Sancak | - | - | Yapay Zekâ |\n| 4 | D işi | Emre Sancak | - | - | - |\n")
    yeni = not_araclari.kisiye_gore_yenile(md)
    assert "## Aksiyonlar" in yeni and "## Kişiye göre iş listesi" in yeni
    assert "- Selin Kara (2): A işi [konu: Veri Ambarı]; B işi" in yeni
    assert "C işi [konu: Yapay Zekâ]" in yeni
    assert [a["konu"] for a in not_araclari.aksiyonlari_ayikla(yeni)] == ["Veri Ambarı", None, "Yapay Zekâ", None]
    # ikinci kez calistirmak bolumu cogaltmaz
    assert not_araclari.kisiye_gore_yenile(yeni).count("## Kişiye göre iş listesi") == 1


def test_jira_labels_konu_slug(tmp_path):
    aks = not_araclari.aksiyonlari_ayikla(NOT_ALTI_SUTUN)
    yol = tmp_path / "a.csv"
    assert not_araclari.jira_csv(aks, str(yol), "Haftalık (2026-10-12)") == 3
    import csv
    with open(yol, encoding="utf-8-sig", newline="") as f:
        satirlar = list(csv.reader(f))
    assert satirlar[0][-1] == "Labels" and satirlar[0][:5] == not_araclari.JIRA_BASLIKLAR
    assert [s[-1] for s in satirlar[1:]] == ["veri-ambari", "yapay-zeka", ""]       # ASCII, bosluk -> '-', bos serbest
    assert all(" " not in s[-1] for s in satirlar[1:])
    assert satirlar[1][1] == "Selin Kara" and satirlar[2][1] == ""                  # sahip kurali degismez
    assert not_araclari.konu_etiketi("Satın Alma / Tedarik") == "satin-alma-tedarik"


def test_jira_konu_yoksa_eski_baslik(tmp_path):
    aks = not_araclari.aksiyonlari_ayikla(NOT_ESKI_BES_SUTUN)
    yol = tmp_path / "b.csv"
    not_araclari.jira_csv(aks, str(yol))
    ilk = yol.read_text(encoding="utf-8-sig").splitlines()[0]
    assert ilk == "Summary,Assignee,Due Date,Description,Issue Type"                # konu yoksa sutun eklenmez


def test_word_baslik_seviyeleri_ve_damga(tmp_path):
    md = ("# Durum\n\n## Konular\n### Veri Ambarı\nDurum: İyi. Risk: Yok.\n#### Kararlar\n- Şema onaylandı\n"
          "#### Açık sorular\n- Bütçe kimde? [konu: Veri Ambarı] ⏱10:02:00\n\n## Aksiyonlar\n"
          "| # | Madde | Sorumlu | Tarih | Kaynak | Konu |\n|---|---|---|---|---|---|\n"
          "| 1 | Şema | Selin Kara | - | - | Veri Ambarı |\n")
    yol = not_araclari.word_kaydet(md, str(tmp_path / "d.docx"))
    import docx
    belge = docx.Document(yol)
    seviye = {p.text: p.style.name for p in belge.paragraphs if p.style.name.startswith("Heading")}
    assert seviye["Konular"] == "Heading 1" and seviye["Veri Ambarı"] == "Heading 3"
    assert seviye["Kararlar"] == "Heading 4" and seviye["Açık sorular"] == "Heading 4"
    metin = "\n".join(p.text for p in belge.paragraphs)
    assert "[konu:" not in metin and "Bütçe kimde? (Konu: Veri Ambarı) ⏱10:02:00" in metin
    assert len(belge.tables) == 1 and [c.text for c in belge.tables[0].rows[0].cells][-1] == "Konu"


def test_epostada_konu_gorunur():
    e = not_araclari.kisiye_ozel_epostalar(NOT_ALTI_SUTUN, "Haftalık", "2026-10-12")
    assert [x[0] for x in e] == ["Selin Kara"]
    assert "- Şemayı hazırlamak (tarih: 2026-10-20) (konu: Veri Ambarı)" in e[0][2]


def test_ara_konu_damgasini_gostermez(tmp_path):
    k = _klasor(tmp_path)
    (k / "not.md").write_text(NOT_ALTI_SUTUN, encoding="utf-8")
    sonuc = [x for x in not_araclari.ara("bütçe onaylandı", kok=str(tmp_path / "toplantilar")) if x["tur"] == "not"]
    assert sonuc and "[konu:" not in sonuc[0]["metin"] and sonuc[0]["ts"] == "10:02:00"


def test_sablon_canli_satirlari_konu_durumlari():
    satirlar = not_araclari.sablon_canli_satirlari(
        {"konu_durumlari": [{"konu": "Veri Ambarı", "durum": "Şema tamam", "risk": "Takvim sıkışık"},
                            {"konu": "", "durum": "Genel akış iyi", "risk": ""}]})
    assert satirlar == [("Konu durumu", "Veri Ambarı: Şema tamam (risk: Takvim sıkışık)"),
                        ("Konu durumu", "Genel: Genel akış iyi")]
    assert not_araclari.sablon_canli_satirlari({"konu_durumlari": []}) == []
    assert not_araclari.sablon_canli_satirlari({}) == []                        # diger sablonlar etkilenmez


def test_konu_damgasi_tek_regex():
    from metin import KONU_DAMGASI
    assert KONU_DAMGASI.sub("", "Karar [konu: Veri Ambarı] ⏱10:02:00") == "Karar ⏱10:02:00"
    assert KONU_DAMGASI.search("x [konu:  Yapay Zekâ ]").group(1).strip() == "Yapay Zekâ"

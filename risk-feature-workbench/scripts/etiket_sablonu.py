"""Web sonuç Excel'inden elle etiketleme şablonu üretir.

Kullanım:
  py -3.9 scripts/etiket_sablonu.py --rapor sonuclar.xlsx --cikti etiketleme.xlsx [--ornek 60] [--tohum 7]

Çıktı "Etiket" sayfası scripts/evaluate.py gerçek moduyla (--label-column "Beklenen kolon")
doğrudan okunur. Yalnız metadata taşır; yeni bağımlılık yok (stdlib + openpyxl).
"""

import argparse
import random
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

SHEET_IN = "Sonuçlar"
BASE = ["Kategori", "Alan Adı", "Alan tipi", "Açıklama"]
REPORTED = "Halihazırda bildiriliyor mu"
STATUS = "Satır durumu"
N_CAND = 3
LABEL = "Beklenen kolon"
NOTE = "Not"

HELP = [
    "Bu şablonda her satır için sistemin önerdiği en çok 3 aday kolon vardır.",
    "Beklenen kolon: doğru kolon(lar)ı SCHEMA.TABLO.KOLON biçiminde yazın; aday hücresinden kopyalayabilirsiniz.",
    "Aday hücresi üç parçalı değilse SCHEMA.TABLO.KOLON biçimine tamamlayın; TABLO.KOLON da kabul edilir (tek şemadaysa).",
    "Birden çok doğru kolon varsa ';' ile ayırın (örn. A.T1.K1;A.T2.K2).",
    "Veritabanında karşılığı yoksa YOK yazın.",
    "Emin değilseniz hücreyi BOŞ bırakın; o satır ölçüme girmez.",
    "Doğru kolon aday listesinde yoksa ama biliyorsanız onu yazın.",
    "Not sütunu serbesttir; ölçüme girmez. Başlık satırını ve sütun adlarını değiştirmeyin.",
    "Dosyada yalnız metadata vardır; parola, bağlantı adresi veya müşteri verisi eklemeyin.",
]


class UserError(Exception):
    pass


def norm(value):
    text = str(value or "").replace("İ", "i").replace("I", "ı").lower().replace("ı", "i")
    text = "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))
    return " ".join(text.split())


def cell_text(value):
    return "" if value is None else str(value)


def read_report(path):
    path = Path(path)
    if not path.is_file():
        raise UserError("Rapor dosyası bulunamadı: " + str(path))
    if path.suffix.lower() != ".xlsx":
        raise UserError("Rapor .xlsx olmalı: " + str(path))
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        raise UserError("Rapor dosyası okunamadı; geçerli bir .xlsx olduğunu ve açık olmadığını kontrol edin.") from None
    try:
        if SHEET_IN not in wb.sheetnames:
            raise UserError("Raporda '" + SHEET_IN + "' sayfası yok; web arayüzünden indirilen sonuç dosyasını verin.")
        rows = list(wb[SHEET_IN].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise UserError("'" + SHEET_IN + "' sayfası boş.")
    headers = [cell_text(h) for h in rows[0]]
    index = {}
    for i, h in enumerate(headers):
        index.setdefault(norm(h), i)
    if norm("Alan Adı") not in index:
        raise UserError("'" + SHEET_IN + "' sayfasında 'Alan Adı' başlığı yok.")
    data = [r for r in rows[1:] if any(v not in (None, "") for v in r)]
    return index, data


def pick(row, index, header):
    i = index.get(norm(header))
    return cell_text(row[i]) if i is not None and i < len(row) else ""


def candidate_path(text):
    # excelio "Aday N kaynak": SCHEMA.OBJECT.COLUMN (eksik parçalar atlanmış olabilir); çözülemezse olduğu gibi.
    text = text.strip()
    parts = [p.strip() for p in text.split(".")]
    if len(parts) == 3 and all(parts) and "\n" not in text:
        return ".".join(parts)
    return text


def stratified_sample(keys, k, seed):
    """keys: satır başına tabaka anahtarı. Seçilen indeksleri (sıralı) döndürür."""
    total = len(keys)
    if k is None or k >= total:
        return list(range(total))
    groups = {}
    for i, key in enumerate(keys):
        groups.setdefault(key, []).append(i)
    names = sorted(groups)
    quota = {g: max(1, int(len(groups[g]) * k / total)) for g in names}
    # Toplam k'yı yakala: önce kalan en büyük kesirlilere ekle, taşarsa en büyük tabakalardan al (min 1).
    while sum(quota.values()) < k:
        room = [g for g in names if quota[g] < len(groups[g])]
        if not room:
            break
        g = max(room, key=lambda x: (len(groups[x]) * k / total - quota[x], len(groups[x]), x))
        quota[g] += 1
    while sum(quota.values()) > k:
        room = [g for g in names if quota[g] > 1]
        if not room:
            break
        g = max(room, key=lambda x: (quota[x] - len(groups[x]) * k / total, quota[x], x))
        quota[g] -= 1
    rng = random.Random(seed)
    chosen = []
    for g in names:
        chosen += rng.sample(groups[g], quota[g])
    return sorted(chosen)


def build(index, data, sample, seed):
    has_reported = norm(REPORTED) in index
    headers = list(BASE) + ([REPORTED] if has_reported else []) + [STATUS]
    for n in range(1, N_CAND + 1):
        headers += ["Aday %d" % n, "Aday %d açıklama" % n, "Aday %d iş anlamı" % n]
    headers += [LABEL, NOTE]
    keys = [(pick(r, index, STATUS), pick(r, index, REPORTED) if has_reported else "") for r in data]
    chosen = stratified_sample(keys, sample, seed)
    out = []
    for i in chosen:
        r = data[i]
        row = [pick(r, index, h) for h in BASE]
        if has_reported:
            row.append(pick(r, index, REPORTED))
        row.append(pick(r, index, STATUS))
        for n in range(1, N_CAND + 1):
            row += [candidate_path(pick(r, index, "Aday %d kaynak" % n)),
                    pick(r, index, "Aday %d kolon açıklaması" % n),
                    pick(r, index, "Aday %d iş anlamı" % n)]
        row += ["", ""]
        out.append(row)
    return headers, out


def write(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Etiket"
    wrap = Alignment(vertical="top", wrap_text=True)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"  # "=" ile başlayan metin formül olmasın
            c.alignment = wrap
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {"Kategori": 18, "Alan Adı": 28, "Alan tipi": 14, "Açıklama": 48, REPORTED: 16, STATUS: 24, LABEL: 36, NOTE: 30}
    for i, h in enumerate(headers, 1):
        if h in widths:
            w = widths[h]
        elif h.endswith("açıklama"):
            w = 40
        elif h.endswith("iş anlamı"):
            w = 22
        else:
            w = 34
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "C2"  # başlık satırı + Kategori/Alan Adı sütunları sabit
    help_ws = wb.create_sheet("Nasıl")
    for line in HELP:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 110
    for row in help_ws.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(target)
    except OSError:
        raise UserError("Çıktı yazılamadı (dosya açık ya da yazma izni yok): " + str(target)) from None


def main(argv=None):
    p = argparse.ArgumentParser(description="Sonuç Excel'inden etiketleme şablonu üretir.")
    p.add_argument("--rapor", required=True, help="web arayüzünden indirilen sonuç .xlsx")
    p.add_argument("--cikti", required=True, help="üretilecek etiketleme .xlsx")
    p.add_argument("--ornek", type=int, help="tabakalı örnek satır sayısı (verilmezse tümü)")
    p.add_argument("--tohum", type=int, default=7, help="örnekleme tohumu (varsayılan 7)")
    args = p.parse_args(argv)
    try:
        if args.ornek is not None and args.ornek < 1:
            raise UserError("--ornek en az 1 olmalı.")
        if Path(args.rapor).resolve() == Path(args.cikti).resolve():
            raise UserError("Çıktı dosyası rapor dosyasının üzerine yazılamaz.")
        index, data = read_report(args.rapor)
        if not data:
            raise UserError("Raporda satır yok.")
        headers, rows = build(index, data, args.ornek, args.tohum)
        write(args.cikti, headers, rows)
    except UserError as exc:
        print("Hata: " + str(exc), file=sys.stderr)
        return 2
    print("Etiket şablonu yazıldı: %s (%d satır, rapordaki %d satırdan)" % (args.cikti, len(rows), len(data)))
    return 0


if __name__ == "__main__":
    sys.exit(main())

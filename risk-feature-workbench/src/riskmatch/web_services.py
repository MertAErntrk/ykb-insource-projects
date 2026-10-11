"""Live-only web forms, bounded workbook inspection and connection probes."""
from copy import deepcopy
import json
import math
import re
import ssl
import urllib.request
from urllib.parse import urlsplit

from .adapters._http import _NoRedirect, post_json, render_template, response_path, validate_endpoint
from .oracle_diagnostics import diagnose_oracle
from .adapters.qwen import LLMAdapter
from .config import load_config
from .errors import ConfigError, InputError, IntegrationError
from .excelio import _load, norm, read_requests

FIELDS = {"name":"Alan Adı", "description":"Açıklama", "category":"Kategori",
          "data_type":"Alan tipi", "reported":"Halihazırda bildiriliyor mu", "row_id":"Talep Kimliği"}
ALIASES = {"name":["Alan Adı","Değişken Adı","Kolon Adı","Variable Name","Column Name"],
           "description":["Açıklama","Tanım","Description"], "category":["Kategori","Category"],
           "data_type":["Alan tipi","Veri tipi","Data Type"],
           "reported":["Halihazırda bildiriliyor mu","Bildiriliyor mu"], "row_id":["Talep Kimliği"]}
PROFILE_KEYS = {
    "oracle": {"dsn","user","source_database","environment","container","schemas","wallet_dir","max_records"},
    "llm": {"base_url","endpoint","model","model_version","contract","request_template","response_path","models_endpoint","ca_file"},
    "embedding": {"base_url","endpoint","model","model_version","dimensions","request_template","response_path","models_endpoint","ca_file"},
}


def clean_profile(raw):
    if not isinstance(raw,dict):raise ConfigError("Ayarlar JSON nesnesi olmalı.")
    if raw.get("mode","real") != "real":raise ConfigError("Web uygulaması yalnız gerçek kurum bağlantılarıyla çalışır.")
    if not isinstance(raw.get("use_models",False),bool):raise ConfigError("Model seçimi geçersiz.")
    result={"mode":"real","use_models":raw.get("use_models",False)}
    for section,keys in PROFILE_KEYS.items():
        values=raw.get(section,{})
        if not isinstance(values,dict):raise ConfigError("Bağlantı bölümleri nesne olmalı.")
        result[section]={k:deepcopy(values[k]) for k in keys if k in values}
    # Templates must not become a back door for saving credentials.
    def inspect(value):
        if isinstance(value,dict):
            for key,item in value.items():
                if str(key).lower() in {"password","api_key","token","authorization","secret","wallet_password"}:
                    raise ConfigError("İstek şablonuna parola veya anahtar yazmayın; parola alanını kullanın.")
                inspect(item)
        elif isinstance(value,list):
            for item in value:inspect(item)
    inspect(result)
    if len(json.dumps(result,ensure_ascii=False))>50000:raise ConfigError("Ayar dosyası çok büyük.")
    return result


def text(value,label,required=False,limit=4096):
    if value is None:value=""
    if not isinstance(value,str) or len(value)>limit:raise ConfigError(label+" geçersiz.")
    value=value.strip()
    if required and not value:raise ConfigError(label+" gerekli.")
    return value


def normalize_dsn(value):
    value=text(value,"Oracle bağlantı metni",True,8000)
    # SQL Developer / tnsnames.ora alias assignment is not part of a connect descriptor.
    match=re.match(r"^\s*[A-Za-z0-9_.-]+\s*=\s*(\(DESCRIPTION\s*=.*)$",value,re.I|re.S)
    return match.group(1).strip() if match else value


def oracle_inputs(raw,probe=False):
    value=raw.get("oracle",{})
    if not isinstance(value,dict):raise ConfigError("Oracle ayarları gerekli.")
    schemas=value.get("schemas",[])
    if isinstance(schemas,str):schemas=[s.strip() for s in re.split(r"[,;\n]",schemas) if s.strip()]
    if not isinstance(schemas,list) or len(schemas)>30:raise ConfigError("En fazla 30 açık şema seçin.")
    limit=value.get("max_records",10000)
    if isinstance(limit,bool) or not isinstance(limit,int) or not 1<=limit<=10000:raise ConfigError("Katalog sınırı 1–10000 olmalı.")
    config={"source_database":text(value.get("source_database"),"Kaynak adı") or "ORACLE",
            "environment":text(value.get("environment"),"Ortam") or "KURUM",
            "container":"PROBE" if probe else text(value.get("container"),"Container",True),
            "schemas":["PROBE"] if probe else schemas,"user_env":"RISK_ORACLE_USER",
            "password_env":"RISK_ORACLE_PASSWORD","dsn_env":"RISK_ORACLE_DSN","max_records":limit}
    credentials={"dsn":normalize_dsn(value.get("dsn")),"user":text(value.get("user"),"Oracle kullanıcı adı",True,256),
                 "password":value.get("password","")}
    if not isinstance(credentials["password"],str) or not credentials["password"]:raise ConfigError("Oracle parolası gerekli.")
    wallet=text(value.get("wallet_dir"),"Wallet dizini")
    if wallet:
        config.update(wallet_dir_env="RISK_ORACLE_WALLET_DIR",wallet_password_env="RISK_ORACLE_WALLET_PASSWORD")
        credentials.update(wallet_dir=wallet,wallet_password=value.get("wallet_password",""))
    return config,credentials


def model_config(raw,name):
    value=raw.get(name,{})
    label="LLM" if name=="llm" else "Embedding"
    if not isinstance(value,dict):raise ConfigError(label+" ayarları gerekli.")
    endpoint=text(value.get("endpoint"),label+" tam çağrı adresi",True)
    parts=urlsplit(endpoint)
    if parts.query:raise ConfigError(label+" adresine anahtar veya sorgu parametresi eklemeyin.")
    config={"endpoint":endpoint,"allowed_hosts":[parts.hostname],
            "model":text(value.get("model"),label+" model kimliği",True),
            "model_version":text(value.get("model_version"),label+" model sürümü",True),
            "timeout_seconds":30,"max_retries":1}
    if name=="llm":config.update(contract=value.get("contract",""),max_candidates=8,max_chars=1800,max_output_tokens=1200)
    else:config.update(dimensions=value.get("dimensions",0),max_batch=16)
    for key in ("request_template","response_path","ca_file"):
        if value.get(key):config[key]=deepcopy(value[key])
    validate_endpoint(config)
    return config


def build_run(raw):
    profile=clean_profile(raw)
    oracle,credentials=oracle_inputs(raw)
    config={"mode":"real","allow_real_connections":True,"use_models":profile["use_models"],"oracle":oracle}
    secrets={"oracle":credentials}
    if profile["use_models"]:
        for name in ("llm","embedding"):
            config[name]=model_config(raw,name)
            key=raw[name].get("api_key","")
            if not isinstance(key,str) or len(key)>8192:raise ConfigError("Servis anahtarı geçersiz.")
            secrets[name]=key
    return load_config(mode="real",data=config),secrets


def probe_oracle(raw):
    config,credentials=oracle_inputs(raw,probe=True)
    return diagnose_oracle(config,credentials)


def probe_model(raw,name):
    config=model_config(raw,name)
    key=raw[name].get("api_key","")
    if name=="llm":
        adapter=LLMAdapter(config,api_key=key)
        adapter.evaluate({"name":"CONNECTION_CHECK","description":"Bağlantı doğrulama"},
                         [{"catalog_id":"connection-check","column":"CONNECTION_CHECK","column_comment":"Bağlantı doğrulama"}])
        return {"ok":True,"message":"LLM çağrısı ve yanıt sözleşmesi doğrulandı."}
    template=config.get("request_template")
    if not isinstance(template,dict) or "{texts}" not in json.dumps(template):raise ConfigError("Embedding istek şablonunda {texts} alanı gerekli.")
    result=post_json(config,render_template(template,{"texts":["Bağlantı doğrulama"],"model":config["model"],"model_version":config["model_version"]}),api_key=key)
    vectors=response_path(result,config.get("response_path"))
    if not isinstance(vectors,list) or len(vectors)!=1 or not isinstance(vectors[0],list):raise IntegrationError("Embedding yanıtı tek vektör içermeli; yanıt yolunu kontrol edin.")
    vector=vectors[0]
    if not 1<=len(vector)<=8192 or any(isinstance(v,bool) or not isinstance(v,(float,int)) or not math.isfinite(v) for v in vector):raise IntegrationError("Embedding vektör biçimi geçersiz.")
    return {"ok":True,"dimensions":len(vector),"message":"Embedding çağrısı doğrulandı; boyut yanıttan okundu."}


def list_models(raw,name):
    value=raw.get(name,{})
    endpoint=text(value.get("models_endpoint"),"Model listesi tam adresi",True)
    config={"endpoint":endpoint,"allowed_hosts":[urlsplit(endpoint).hostname]}
    validate_endpoint(config)
    headers={"Accept":"application/json"}
    if value.get("api_key"):headers["Authorization"]="Bearer "+value["api_key"]
    try:
        context=ssl.create_default_context(cafile=value.get("ca_file") or None)
        opener=urllib.request.build_opener(_NoRedirect(),urllib.request.HTTPSHandler(context=context))
        with opener.open(urllib.request.Request(endpoint,headers=headers),timeout=15) as response:
            body=response.read(1000001)
        if len(body)>1000000:raise ValueError()
        data=json.loads(body)
        ids=[item["id"] for item in data["data"] if isinstance(item,dict) and isinstance(item.get("id"),str)]
        return {"models":ids[:100]}
    except Exception:raise IntegrationError("Model listesi alınamadı. Adres, erişim ve OpenAI liste sözleşmesini kontrol edin; model kimliğini elle de girebilirsiniz.") from None


def inspect_workbook(path):
    wb=_load(path)
    sheets=[]
    try:
        for sheet in wb:
            if sheet.max_row>20000 or sheet.max_column>300:raise InputError("Excel sınırı: sayfa başına 20.000 satır ve 300 sütun.")
            choices=[]
            for row in sheet.iter_rows(max_row=min(30,sheet.max_row)):
                values=[str(c.value or "") for c in row]
                score=sum(any(norm(v) in {norm(a) for a in names} for v in values) for names in ALIASES.values())
                if any(values):choices.append((score,row[0].row,values))
            if not choices:continue
            _,header,labels=max(choices,key=lambda entry:(entry[0],-entry[1]))
            suggested={key:next((i+1 for i,label in enumerate(labels) if norm(label) in {norm(a) for a in aliases}),None) for key,aliases in ALIASES.items()}
            headers=[{"index":i+1,"label":label or "Sütun "+str(i+1)} for i,label in enumerate(labels)]
            preview=[[str(c.value or "")[:500] for c in row] for row in sheet.iter_rows(min_row=header+1,max_row=min(header+3,sheet.max_row))]
            sheets.append({"name":sheet.title,"header_row":header,"rows":max(0,sheet.max_row-header),"headers":headers,"suggested":suggested,"preview":preview})
        if not sheets:raise InputError("Çalışma kitabında dolu sayfa bulunamadı.")
        return sheets
    finally:wb.close()


def validate_mapping(path,mapping):
    if not isinstance(mapping,dict) or not mapping:raise InputError("En az bir Excel sayfası seçin.")
    wb=_load(path)
    try:
        for name,settings in mapping.items():
            if name not in wb.sheetnames or not isinstance(settings,dict):raise InputError("Excel sayfa seçimi geçersiz.")
            row=settings.get("header_row")
            if isinstance(row,bool) or not isinstance(row,int) or not 1<=row<=min(wb[name].max_row,30):raise InputError("Başlık satırı 1–30 aralığında olmalı.")
            columns=settings.get("columns")
            if not isinstance(columns,dict) or set(columns)-set(FIELDS):raise InputError("Sütun eşlemesi geçersiz.")
            if not columns.get("name") or not columns.get("description"):raise InputError("Alan adı ve açıklama sütunlarını seçin.")
            indexes=[i for i in columns.values() if i is not None]
            if any(isinstance(i,bool) or not isinstance(i,int) or not 1<=i<=wb[name].max_column for i in indexes):raise InputError("Sütun numarası geçersiz.")
            if len(set(indexes))!=len(indexes):raise InputError("Aynı sütun iki farklı alana atanamaz.")
    finally:wb.close()
    requests=read_requests(path,mapping)
    if len(requests)>2000:raise InputError("Bir analizde en fazla 2.000 değişken seçin; Excel'i bölerek çalıştırın.")
    return requests


def automatic_mapping(path):
    sheets=inspect_workbook(path)
    result={s["name"]:{"header_row":s["header_row"],"columns":s["suggested"]}
            for s in sheets if s["suggested"]["name"] and s["suggested"]["description"]}
    if not result:raise InputError("Alan adı ve açıklama başlıkları tanınmadı. Sütun eşleştirmesini açıp seçin.")
    return result


def write_manual_request(path, values):
    from openpyxl import Workbook
    if not isinstance(values,list) or not 1<=len(values)<=100:
        raise InputError("Ekrandan 1–100 değişken ekleyebilirsiniz. Daha fazlası için Excel kullanın.")
    rows=[]
    for index,value in enumerate(values,1):
        if not isinstance(value,dict):raise InputError("Değişken satırı geçersiz.")
        name=text(value.get("name"),str(index)+". değişkenin alan adı",True,256)
        description=text(value.get("description"),str(index)+". değişkenin açıklaması",True,8000)
        rows.append([name,description])
    wb=Workbook();sheet=wb.active;sheet.title="Ekrandan eklenenler"
    sheet.append(["Alan Adı","Açıklama"])
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:cell.data_type="s"
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path);wb.close()
    return {"Ekrandan eklenenler":{"header_row":1,"columns":{"name":1,"description":2}}}

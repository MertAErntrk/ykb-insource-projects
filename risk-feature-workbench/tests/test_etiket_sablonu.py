import contextlib
import io
import sys
import tempfile
import types
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "src"))

import etiket_sablonu  # noqa: E402
from riskmatch.excelio import read_requests, write_results  # noqa: E402

STATUSES = ["found", "found", "found", "found", "weak_only", "weak_only", "none_in_scope", "none_in_scope"]


def make_report(folder, count=24):
    src = Path(folder) / "talep.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "DegiskenSozlugu"
    ws.append(["Kategori", "Alan Adı", "Alan tipi", "Açıklama", "Halihazırda bildiriliyor mu"])
    for i in range(count):
        ws.append(["Kat%d" % (i % 3), "ALAN_%d" % i, "NUMBER", "Açıklama %d" % i, "Evet" if i % 2 else "Hayır"])
    wb.save(src)
    requests = read_requests(src)
    results = []
    for i, req in enumerate(requests):
        status = STATUSES[i % len(STATUSES)]
        cands = []
        if status != "none_in_scope":
            for n in range(2):
                cands.append({"catalog_id": "c%d_%d" % (i, n), "source": {"schema": "EDW", "object": "TAB%d" % n, "column": "COL_%d" % i, "column_comment": "kolon %d" % n},
                              "name_match": "exact", "meaning": "likely", "locations": []})
        results.append({"row_id": req["row_id"], "row_status": status, "candidates": cands})
    out = Path(folder) / "sonuclar.xlsx"
    write_results(src, out, requests, results, {"Sürüm": "test"})
    return out


class EtiketSablonuTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.dir = self._tmp.name
        self.report = make_report(self.dir)

    def run_cli(self, *extra, rapor=None, cikti="etiket.xlsx"):
        out = Path(self.dir) / cikti
        err = io.StringIO()
        with contextlib.redirect_stderr(err), contextlib.redirect_stdout(io.StringIO()):
            code = etiket_sablonu.main(["--rapor", str(rapor or self.report), "--cikti", str(out), *extra])
        return code, out, err.getvalue()

    def rows(self, path):
        ws = load_workbook(path)["Etiket"]
        return [[c for c in r] for r in ws.iter_rows(values_only=True)]

    def test_headers_and_candidate_paths(self):
        code, out, _ = self.run_cli()
        self.assertEqual(code, 0)
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Etiket", "Nasıl"])
        ws = wb["Etiket"]
        expected = ["Kategori", "Alan Adı", "Alan tipi", "Açıklama", "Halihazırda bildiriliyor mu", "Satır durumu",
                    "Aday 1", "Aday 1 açıklama", "Aday 1 iş anlamı", "Aday 2", "Aday 2 açıklama", "Aday 2 iş anlamı",
                    "Aday 3", "Aday 3 açıklama", "Aday 3 iş anlamı", "Beklenen kolon", "Not"]
        self.assertEqual([c.value for c in ws[1]], expected)
        self.assertTrue(all(c.font.bold for c in ws[1]))
        self.assertEqual(ws.freeze_panes, "C2")
        rows = self.rows(out)
        self.assertEqual(len(rows), 25)
        first = dict(zip(expected, rows[1]))
        self.assertEqual(first["Aday 1"], "EDW.TAB0.COL_0")
        self.assertEqual(first["Aday 2"], "EDW.TAB1.COL_0")
        self.assertIn(first["Aday 3"], ("", None))
        self.assertEqual(first["Aday 1 açıklama"], "kolon 0")
        self.assertEqual(first["Aday 1 iş anlamı"], "Olası uyum")
        self.assertIn(first["Beklenen kolon"], ("", None))
        self.assertEqual(len(wb["Nasıl"]["A"]), len(etiket_sablonu.HELP))

    def test_sample_deterministic_and_stratified(self):
        code, a, _ = self.run_cli("--ornek", "8", "--tohum", "3", cikti="a.xlsx")
        code2, b, _ = self.run_cli("--ornek", "8", "--tohum", "3", cikti="b.xlsx")
        self.assertEqual((code, code2), (0, 0))
        ra, rb = self.rows(a), self.rows(b)
        self.assertEqual(ra, rb)
        self.assertEqual(len(ra), 9)
        header = ra[0]
        si, ri = header.index("Satır durumu"), header.index("Halihazırda bildiriliyor mu")
        strata = {(r[si], r[ri]) for r in ra[1:]}
        full = {(r[si], r[ri]) for r in self.rows(self.run_cli(cikti="f.xlsx")[1])[1:]}
        self.assertEqual(strata, full)  # her tabakadan en az 1
        _, c, _ = self.run_cli("--ornek", "8", "--tohum", "99", cikti="c.xlsx")
        self.assertEqual(len(self.rows(c)), 9)

    def test_readable_by_read_requests(self):
        _, out, _ = self.run_cli()
        wb = load_workbook(out)
        ws = wb["Etiket"]
        col = [c.value for c in ws[1]].index("Beklenen kolon") + 1
        ws.cell(2, col, "EDW.TAB0.COL_0;YOK")
        wb.save(out)
        reqs = [r for r in read_requests(out) if r["sheet"] == "Etiket"]
        self.assertEqual(len(reqs), 24)
        self.assertFalse([r for r in reqs if r["input_errors"]])
        self.assertEqual(reqs[0]["name"], "ALAN_0")
        self.assertEqual(reqs[0]["description"], "Açıklama 0")
        self.assertEqual(reqs[0]["original"]["Beklenen kolon"], "EDW.TAB0.COL_0;YOK")
        self.assertIn("Beklenen kolon", reqs[1]["original"])

    def test_evaluate_reads_template_blank_is_unlabelled(self):
        import evaluate
        _, out, _ = self.run_cli()
        wb = load_workbook(out)
        ws = wb["Etiket"]
        col = [c.value for c in ws[1]].index("Beklenen kolon") + 1
        ws.cell(2, col, "EDW.TAB0.COL_0")
        ws.cell(3, col, "YOK")
        wb.save(out)
        args = types.SimpleNamespace(excel=str(out), column_mapping=None, label_column="Beklenen kolon")
        rows = evaluate.load_rows(args)
        self.assertEqual(len(rows), 24)  # the help sheet is not read as requests
        self.assertEqual(args._atlanan["adet"], 0)
        self.assertEqual([r["etiket"] for r in rows[:3]], ["pozitif", "yok", "etiketsiz"])
        self.assertEqual(sum(1 for r in rows if r["label_valid"]), 2)

    def test_errors_exit_2(self):
        code, _, err = self.run_cli(rapor=Path(self.dir) / "yok.xlsx")
        self.assertEqual(code, 2)
        self.assertIn("bulunamadı", err)
        self.assertNotIn("Traceback", err)
        bad = Path(self.dir) / "bos.xlsx"
        wb = Workbook()
        wb.active.title = "Sonuçlar"
        wb.active.append(["Foo", "Bar"])
        wb.save(bad)
        code, _, err = self.run_cli(rapor=bad)
        self.assertEqual(code, 2)
        self.assertIn("Alan Adı", err)
        wb = Workbook()
        wb.active.title = "Baska"
        wb.save(bad)
        code, _, err = self.run_cli(rapor=bad)
        self.assertEqual(code, 2)
        self.assertNotIn("Traceback", err)

    def test_unparsed_candidate_kept(self):
        self.assertEqual(etiket_sablonu.candidate_path("TAB.COL"), "TAB.COL")
        self.assertEqual(etiket_sablonu.candidate_path(" A.B.C "), "A.B.C")


if __name__ == "__main__":
    unittest.main()

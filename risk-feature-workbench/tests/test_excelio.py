from pathlib import Path
import glob
import tempfile
import time
import unittest
from unittest import mock

from openpyxl import Workbook,load_workbook
from riskmatch.errors import InputError,ReportWriteError
from riskmatch import excelio
from riskmatch.excelio import HEADERS,read_requests,write_results,prepare_requests

class ExcelTests(unittest.TestCase):
    def make(self,path):
        wb=Workbook()
        ws=wb.active
        ws.title="Girdi"
        ws.append(HEADERS)
        ws.append(["Başvuru","A","Numeric","Türkçe<br>bir\niki","Evet","=1+1","3 gün"])
        ws.append(["Başvuru","A","Numeric",None,"Hayır","Kullanıcı görüşü",None])
        ws.append([None]*7)
        ws.append(["Başvuru","Oran","Numeric","=B2/B3",None,None,None])
        wb.create_sheet("İkinci").append(HEADERS)
        wb["İkinci"].append(["Risk","B","Numeric","",None,None,None])
        wb.save(path)
        wb.close()

    def test_preserves_all_values_sheets_order_and_neutralizes_formula(self):
        with tempfile.TemporaryDirectory() as directory:
            src=Path(directory)/"source.xlsx"
            dst=Path(directory)/"out.xlsx"
            self.make(src)
            req=read_requests(src)
            self.assertEqual([r["name"] for r in req],["A","A","Oran","B"])
            self.assertEqual(len({r["row_id"] for r in req}),4)
            self.assertEqual(req[2]["excel_row"],5)
            self.assertEqual(req[2]["description_for_parse"],"A / A")
            rows=[{"row_id":r["row_id"],"physical_name":"unknown","semantic_status":"insufficient","production_status":"insufficient","reasons":[],"errors":[],"candidates":[]} for r in req]
            write_results(src,dst,req,rows,{"Mod":"DEMO"})
            wb=load_workbook(dst,data_only=False)
            self.assertEqual(wb.sheetnames[:3],["Sonuçlar","Adaylar","Çalışma bilgisi"])
            self.assertEqual(wb["Girdi 1 Girdi"]["D2"].value,"Türkçe<br>bir\niki")
            self.assertEqual(wb["Girdi 1 Girdi"]["F2"].value,"=1+1")
            self.assertEqual(wb["Girdi 1 Girdi"]["F2"].data_type,"s")
            self.assertEqual(wb["Girdi 1 Girdi"]["G2"].value,"3 gün")
            self.assertIsNone(wb["Girdi 1 Girdi"]["B4"].value)
            self.assertEqual(wb["Sonuçlar"].max_row,5)
            wb.close()

    def test_ids_persist_through_edit_of_prepared_copy(self):
        with tempfile.TemporaryDirectory() as directory:
            src=Path(directory)/"source.xlsx"
            identified=Path(directory)/"identified.xlsx"
            self.make(src)
            before=read_requests(src)
            prepare_requests(src,identified,before)
            wb=load_workbook(identified)
            wb["Girdi"]["D2"]="Düzeltilmiş açıklama"
            wb.save(identified)
            wb.close()
            after=read_requests(identified)
            self.assertEqual([r["row_id"] for r in before],[r["row_id"] for r in after])

    def test_unrecognized_sheet_and_formula_exposed(self):
        with tempfile.TemporaryDirectory() as directory:
            src=Path(directory)/"in.xlsx"
            self.make(src)
            wb=load_workbook(src)
            wb["Girdi"]["D5"]="=WEBSERVICE(\"http://never.invalid\")"
            wb.create_sheet("Bozuk").append(["Alanları içeriyor ama başlık yok"])
            wb.save(src)
            wb.close()
            req=read_requests(src)
            self.assertTrue(req[2]["input_errors"])
            self.assertTrue(req[-1]["input_errors"])
            with self.assertRaises(InputError):prepare_requests(src,src)

def fake_candidate(i,n,rejected=False):
    c={"catalog_id":f"C{i}-{n}","source":{"schema":"S","object":f"T{n}","column":f"COL{i}","data_type":"NUMBER","column_comment":"Açıklama","table_comment":"Tablo"},"locations":["S.T1","S.T2"],"name_match":"partial","meaning":"likely","differences":[{"dimension":"window","requested":"12 ay","actual":"6 ay","text":"Pencere farklı"}],"evidence":["kanıt 1"],"status":"supported","checks":[{"dimension":"window","status":"conflict","requested":"12 ay","actual":"6 ay","evidence":{}}],"reasons":[],"features":{},"score":0.8}
    if rejected:c["rejected_reason"]="Anlam farkı"
    return c

class ReportWriteTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.dir=Path(self.tmp.name)
        self.src=self.dir/"source.xlsx"
        wb=Workbook();ws=wb.active;ws.title="Girdi";ws.append(HEADERS)
        for i in range(273):
            ws.append(["Kat",f"Alan{i}","Numeric",f"Açıklama {i}",None,None,None])
        wb.save(self.src);wb.close()
        self.req=read_requests(self.src)
    def tearDown(self):
        self.tmp.cleanup()
    def full(self,req):
        return {"row_id":req["row_id"],"physical_name":"unknown","semantic_status":"supported","production_status":"not_evaluated","reasons":[],"errors":[],"dependencies":[],
                "candidates":[fake_candidate(1,n) for n in range(3)],"rejected_candidates":[fake_candidate(2,n,True) for n in range(5)],
                "row_status":"found","coverage":{"catalog_complete":True,"index_coverage":0.97,"semantic_scope":"full"},"model_changed":False}
    def test_speed_and_shape(self):
        results=[self.full(r) for r in self.req]
        dst=self.dir/"out.xlsx"
        start=time.perf_counter()
        write_results(self.src,dst,self.req,results,{"Mod":"DEMO"})
        elapsed=time.perf_counter()-start
        print(f"write_results 273 satır: {elapsed:.2f} s")
        self.assertLess(elapsed,6)
        wb=load_workbook(dst)
        ws=wb["Sonuçlar"]
        header=[c.value for c in ws[1]]
        self.assertEqual(ws.max_row,274)
        self.assertIn("Aday 3 farklar ve kanıt",header)
        self.assertNotIn("Olası şema",header)
        row=[c.value for c in ws[2]]
        self.assertEqual(row[header.index("Satır durumu")],"Aday bulundu")
        self.assertEqual(row[header.index("Aday 1 kaynak")],"S.T0.COL1")
        self.assertEqual(row[header.index("Aday 1 ad eşleşmesi")],"Kısmi ad benzerliği")
        self.assertEqual(row[header.index("Aday 1 iş anlamı")],"Olası uyum")
        self.assertIn("12 ay",row[header.index("Aday 1 farklar ve kanıt")])
        self.assertEqual(wb["Adaylar"].max_row,1+273*8)
        self.assertEqual(wb["Adaylar"]["C2"].value,"Gösterildi")
        self.assertTrue(wb["Adaylar"]["C5"].value.startswith("Reddedildi: Anlam farkı"))
        wb.close()
    def test_old_format_and_subset(self):
        old=[{"row_id":r["row_id"],"physical_name":"unknown","semantic_status":"insufficient","production_status":"insufficient","reasons":[],"errors":[],"candidates":[{"catalog_id":"X","source":{"schema":"S","object":"T","column":"C"},"status":"supported"}]} for r in self.req[:2]]
        dst=self.dir/"old.xlsx"
        write_results(self.src,dst,self.req[:5],old,{"Mod":"DEMO"})
        wb=load_workbook(dst)
        ws=wb["Sonuçlar"]
        header=[c.value for c in ws[1]]
        self.assertEqual(ws.max_row,6)
        self.assertEqual(ws.cell(2,header.index("Aday 1 kaynak")+1).value,"S.T.C")
        self.assertEqual(ws.cell(6,header.index("Satır durumu")+1).value,"Aranmadı")
        wb.close()
    def test_model_changed_warning(self):
        res=[self.full(r) for r in self.req[:1]]
        res[0]["model_changed"]=True
        dst=self.dir/"m.xlsx"
        write_results(self.src,dst,self.req[:1],res,{"Mod":"DEMO"})
        wb=load_workbook(dst)
        header=[c.value for c in wb["Sonuçlar"][1]]
        self.assertIn("Uyarı",wb["Sonuçlar"].cell(2,header.index("Satır durumu")+1).value)
        self.assertIn("Uyarı",[r[0].value for r in wb["Çalışma bilgisi"].iter_rows(min_row=2)])
        wb.close()
    def test_replace_permission_error_keeps_old_report(self):
        dst=self.dir/"keep.xlsx"
        dst.write_bytes(b"OLD REPORT")
        with mock.patch("riskmatch.excelio.os.replace",side_effect=PermissionError(13,"denied")):
            with self.assertRaises(ReportWriteError) as ctx:
                write_results(self.src,dst,self.req[:3],[self.full(r) for r in self.req[:3]],{})
        self.assertIn(ctx.exception.kind,("locked","permission"))
        self.assertEqual(dst.read_bytes(),b"OLD REPORT")
        self.assertEqual(glob.glob(str(self.dir/"*.writing.xlsx")),[])
        self.assertNotIn(str(self.dir),str(ctx.exception))
    def test_failed_verification_keeps_target(self):
        dst=self.dir/"keep.xlsx"
        dst.write_bytes(b"OLD REPORT")
        with mock.patch("riskmatch.excelio._count_rows",return_value=1):
            with self.assertRaises(ReportWriteError) as ctx:
                write_results(self.src,dst,self.req[:3],[self.full(r) for r in self.req[:3]],{})
        self.assertEqual(ctx.exception.kind,"invalid")
        self.assertEqual(dst.read_bytes(),b"OLD REPORT")
        self.assertEqual(glob.glob(str(self.dir/"*.writing.xlsx")),[])
    def test_text_cells_stay_text(self):
        req=[dict(self.req[0],name="=cmd|' /C calc'!A0",original={"Alan Adı":"=cmd|' /C calc'!A0"})]
        dst=self.dir/"f.xlsx"
        write_results(self.src,dst,req,[self.full(req[0])],{})
        wb=load_workbook(dst)
        cell=wb["Sonuçlar"]["D2"]
        self.assertEqual(cell.data_type,"s")
        self.assertTrue(cell.value.startswith("=cmd"))
        wb.close()

class VerdictColumnsTests(unittest.TestCase):
    def results(self,extra):
        with tempfile.TemporaryDirectory() as directory:
            src=Path(directory)/"source.xlsx"
            dst=Path(directory)/"out.xlsx"
            ExcelTests().make(src)
            req=read_requests(src)
            rows=[]
            for i,r in enumerate(req):
                row={"row_id":r["row_id"],"physical_name":"unknown","semantic_status":"insufficient","production_status":"insufficient","reasons":[],"errors":[],"candidates":[]}
                row.update(extra.get(i,{}))
                rows.append(row)
            write_results(src,dst,req,rows,{"Mod":"DEMO"})
            wb=load_workbook(dst)
            ws=wb["Sonuçlar"]
            headers=[c.value for c in ws[1]]
            values=[{h:c.value for h,c in zip(headers,row)} for row in ws.iter_rows(min_row=2)]
            wb.close()
            return headers,values

    def test_verdict_columns_are_appended_after_existing_ones(self):
        direct={"verdict":"dogrudan","verdict_reasons":[{"code":"x","text":"Ad ve anlam uyuşuyor"}],"derivation":None,"tier_trace":[],
                "candidates":[{"catalog_id":"c","source":{"schema":"DM_CAD","object":"T","column":"C"}}]}
        ratio={"verdict":"turetilebilir","verdict_reasons":[{"code":"a","text":"Pay bulundu"},{"code":"b","text":"Aynı tabloda"}],
               "derivation":{"formula":"pay / payda","relationship":"same_object","verdict":"turetilebilir","verdict_reasons":[],
                             "components":[{"role":"numerator","name":"p","window":"1M","verdict":"dogrudan","source":{"schema":"DM","object":"T","column":"P"},"tier":"DM","reasons":[]},
                                           {"role":"denominator","name":"d","window":"1M","verdict":"bulunamadi","source":None,"tier":"DM","reasons":[]}]}}
        headers,values=self.results({0:direct,1:ratio})
        self.assertEqual(headers[-5:],["Uzman kararı","Sonuç","Kaynak","Dayanak","Türetme yöntemi"])
        self.assertEqual(headers[:3],["Talep Kimliği","Kaynak Sayfa","Kaynak Satır"])
        self.assertEqual((values[0]["Sonuç"],values[0]["Kaynak"]),("Doğrudan eşleşti","DM_CAD.T.C"))
        self.assertEqual(values[0]["Dayanak"],"Ad ve anlam uyuşuyor")
        self.assertEqual(values[1]["Sonuç"],"Türetilebilir")
        self.assertEqual(values[1]["Kaynak"],"Pay: DM.T.P; Payda: aranan katmanlarda bulunamadı")
        self.assertEqual(values[1]["Dayanak"],"Pay bulundu; Aynı tabloda")
        self.assertEqual(values[1]["Türetme yöntemi"],"pay / payda (aynı tablo)")

    def test_old_results_without_verdict_leave_cells_empty(self):
        _,values=self.results({})
        for row in values:
            for key in ("Sonuç","Kaynak","Dayanak","Türetme yöntemi"):
                self.assertIn(row[key],(None,""))

if __name__=="__main__":unittest.main()

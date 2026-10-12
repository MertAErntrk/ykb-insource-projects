"""Workspace flow: upload once, search selected rows in groups, report only after a safe write."""
import copy
import io
import os
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from riskmatch.web import create_app
from riskmatch import excelio
from riskmatch.errors import AnalysisCancelled, ReportWriteError
from riskmatch.pipeline import MODEL_CHANGED_REASON

EMBEDDING={"use_models":True,"llm_review":False,"embedding":{"endpoint":"https://embed.invalid/embed","model":"embed-id",
           "model_version":"v1","request_template":{"inputs":"{texts}"},"response_path":"$","dimensions":2,"api_key":"EMBED_SECRET"}}


def workbook(count=10):
    wb=Workbook();ws=wb.active;ws.title="Değişkenler"
    ws.append(["Alan Adı","Açıklama","Kategori"])
    for n in range(count):ws.append(["VAR_"+str(n),"Açıklama "+str(n),"Risk"])
    data=io.BytesIO();wb.save(data);wb.close();data.seek(0);return data


class FakeIndex:
    coverage=0.5
    def stored_coverage(self):return self.coverage
    # Tablo belgeleri de sayılır: kolonlar tam olsa da tablo eksikse indeks tamam değildir.
    tables_missing=0
    def queue_counts(self):return (0 if self.coverage>=1.0 else 5)+self.tables_missing,12


class FakeIndexer:
    def __init__(self):self.state="idle";self.index=FakeIndex();self.stops=0
    def start(self):self.state="running"
    def stop(self):self.state="stopped";self.stops+=1
    def status(self):return {"state":self.state,"done":5,"total":10,"failures":0,"rate_per_s":1.0,"eta_s":5,"message":""}


class FakeResources:
    def __init__(self,root):self.root=root;self.calls=[];self.closed=False;self.last=None;self.closes=[];self.lock=None
    def catalog(self,adapter,key,*,refresh=False,progress=None):
        self.calls.append(("catalog",key,refresh));return {"records":[{"column":"A"},{"column":"B"}]},{"source":"oracle" if refresh else "cache"}
    def embedding(self,adapter):self.calls.append(("embedding",));return "ctx"
    def indexer(self,snapshot,ctx,settings):self.last=FakeIndexer();return self.last
    def close(self):
        # Records whether the app guard was free: a non-reentrant lock held by this thread would time out.
        free=self.lock.acquire(timeout=2) if self.lock else True
        if free and self.lock:self.lock.release()
        self.closed=True;self.closes.append(free)


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.calls=[];self.seen_progress=[];self.resources=[]
        self.gate=None;self.ready=threading.Event();self.cancel_group=None;self.before_group=None;self.infos={}
        def factory(root):
            value=FakeResources(root);value.lock=self.app.extensions["riskmatch"]["guard"];self.resources.append(value);return value
        self.app=create_app(self.root,analyzer=self.analyze,resources_factory=factory)
        self.client=self.app.test_client();self.client.get("/")
        with self.client.session_transaction() as s:self.headers={"X-CSRF-Token":s["csrf"]}
        with patch("riskmatch.web.services.probe_oracle",return_value={"ok":True,"db_name":"EDWLIVE","container":"PDB1","schemas":["DM_CAD"]}):
            self.assertEqual(self.client.post("/api/connect",json={"user":"u","password":"CREDENTIAL_SENTINEL"},headers=self.headers).status_code,200)
        response=self.client.post("/api/upload",data={"file":(workbook(),"talepler.xlsx")},headers=self.headers)
        self.assertEqual(response.status_code,200,response.json)
        self.upload_id=response.json["id"]

    def tearDown(self):
        if self.gate:self.gate.set()
        for job in self.app.extensions["riskmatch"]["jobs"].values():job["cancel_event"].set()
        self.app.extensions["riskmatch"]["executor"].shutdown(wait=True)
        self.tmp.cleanup()

    def row(self,row_id):
        return {"row_id":row_id,"row_status":"found","physical_name":"present","semantic_status":"supported",
                "production_status":"not_evaluated","reasons":[],"errors":[],
                "candidates":[{"catalog_id":"c1","status":"supported","name_match":"exact","meaning":"strong",
                               "source":{"schema":"DM_CAD","object":"T","column":"A"},"locations":[]}]}

    def analyze(self,path,output,**kwargs):
        ids=list(kwargs["row_ids"]);self.calls.append({"row_ids":ids,"write_report":kwargs["write_report"],"resources":kwargs["resources"],
                                                       "mapping":kwargs["column_mapping"],
                                                       "require_full_index":kwargs.get("require_full_index")})
        number=len(self.calls)
        kwargs["progress"]({"stage":"Aranıyor","completed":1,"total":len(ids),"phase":"analysis"})
        if self.before_group:self.before_group(number)
        job=next(j for j in self.app.extensions["riskmatch"]["jobs"].values() if j["status"] in ("running","cancelling"))
        self.seen_progress.append({k:job["progress"][k] for k in ("group","groups","completed","total")})
        rows=[]
        for index,row_id in enumerate(ids):
            row=self.row(row_id)
            if self.cancel_group==len(self.calls) and index==1:
                self.ready.set();self.gate.wait(5)
                if kwargs["cancelled"]():
                    error=AnalysisCancelled("durdu");error.partial={"results":rows+[row]};raise error
            kwargs["result_callback"](row);rows.append(row)
        return {"results":rows,"requests":[],"run_info":self.infos.get(number,{"Arama modu":"fast"}),"output":None,"report_error":None}

    def start(self,**extra):
        response=self.client.post("/api/analyze",json={"upload_id":self.upload_id,"schemas":["DM_CAD"],**extra},headers=self.headers)
        self.assertEqual(response.status_code,202,response.json)
        return response.json["id"]

    def finish(self,identity):
        deadline=time.monotonic()+10
        while time.monotonic()<deadline:
            value=self.client.get("/api/jobs/"+identity).json
            if value["status"] not in ("running","cancelling","reporting"):return value
            time.sleep(.01)
        self.fail("Worker did not finish")

    def downloaded(self,identity):
        response=self.client.get("/api/jobs/"+identity+"/download");response.close()
        return response.status_code

    def row_ids(self):
        return [r["row_id"] for r in self.client.get("/api/uploads/"+self.upload_id+"/rows").json["rows"]]

    def test_first_row_then_remaining_rows_without_reupload(self):
        ids=self.row_ids();self.assertEqual(len(ids),10)
        written=[]
        def capture(input_path,output_path,requests,results,run_info):
            written.append((len(requests),len(results),dict(run_info)));return real(input_path,output_path,requests,results,run_info)
        real=excelio.write_results
        with patch("riskmatch.web.write_results",side_effect=capture):
            first=self.finish(self.start(row_ids=ids[:1]))
            self.assertEqual(first["status"],"complete");self.assertEqual(first["result_count"],1)
            rows=self.client.get("/api/uploads/"+self.upload_id+"/rows").json
            self.assertEqual((rows["total"],rows["searched"]),(10,1))
            self.assertTrue(rows["rows"][0]["searched"]);self.assertEqual(rows["rows"][0]["row_status_label"],"Aday bulundu")
            self.assertFalse(rows["rows"][1]["searched"]);self.assertEqual(rows["rows"][1]["row_status"],"not_searched")
            second=self.finish(self.start(row_ids=ids[1:]+ids[1:2]))
        self.assertEqual(second["status"],"complete");self.assertEqual(second["result_count"],9)
        self.assertEqual(self.client.get("/api/uploads/"+self.upload_id+"/rows").json["searched"],10)
        self.assertEqual([w[:2] for w in written],[(10,1),(10,10)])
        self.assertEqual((written[1][2]["Aranan satır"],written[1][2]["Bu işte aranan"]),(10,9))
        self.assertTrue(all(c["write_report"] is False and c["resources"] is self.resources[0] for c in self.calls))
        self.assertEqual(len(self.resources),1)
        candidate=second["results"][0]["candidates"][0]
        self.assertEqual((candidate["name_match_label"],candidate["meaning_label"]),("Birebir ad","Güçlü anlam uyumu"))
        self.assertEqual(self.downloaded(second["id"]),200)

    def test_groups_split_selection_and_progress_counts_whole_job(self):
        job=self.finish(self.start(group_size=4))
        self.assertEqual([len(c["row_ids"]) for c in self.calls],[4,4,2])
        self.assertEqual(self.seen_progress,[{"group":1,"groups":3,"completed":1,"total":10},
            {"group":2,"groups":3,"completed":5,"total":10},{"group":3,"groups":3,"completed":9,"total":10}])
        self.assertEqual(job["progress"]["completed"],10)

    def test_invalid_selection_and_group_size_rejected(self):
        for extra in ({"row_ids":[]},{"row_ids":["UNKNOWN"]},{"row_ids":"x"},{"group_size":0},{"group_size":101},{"group_size":True}):
            response=self.client.post("/api/analyze",json={"upload_id":self.upload_id,"schemas":["DM_CAD"],**extra},headers=self.headers)
            self.assertEqual(response.status_code,400,extra)
            self.assertTrue(response.json["error"])
        self.assertFalse(self.app.extensions["riskmatch"]["jobs"])

    def test_locked_report_keeps_results_and_retry_completes(self):
        real=excelio.write_results;attempts=[]
        def flaky(*args):
            attempts.append(1)
            if len(attempts)==1:raise ReportWriteError("locked")
            return real(*args)
        with patch("riskmatch.web.write_results",side_effect=flaky):
            job=self.finish(self.start())
            self.assertEqual(job["status"],"report_failed");self.assertEqual(job["error_kind"],"locked")
            self.assertEqual(job["report_error"]["kind"],"locked");self.assertFalse(job["report_ready"])
            self.assertEqual(job["result_count"],10)
            self.assertEqual(self.client.get("/api/jobs/"+job["id"]+"/download").status_code,400)
            self.assertEqual(self.client.post("/api/jobs/"+job["id"]+"/report",json={}).status_code,403)
            retry=self.client.post("/api/jobs/"+job["id"]+"/report",json={},headers=self.headers)
        self.assertEqual((retry.json["status"],retry.json["report_ready"]),("complete",True))
        self.assertEqual(self.client.get("/api/jobs/"+job["id"]).json["status"],"complete")
        self.assertEqual(self.downloaded(job["id"]),200)

    def test_complete_never_shown_before_report_written(self):
        gate=threading.Event();entered=threading.Event();real=excelio.write_results
        def slow(*args):
            entered.set();gate.wait(5);return real(*args)
        with patch("riskmatch.web.write_results",side_effect=slow):
            identity=self.start()
            self.assertTrue(entered.wait(5))
            for _ in range(5):
                value=self.client.get("/api/jobs/"+identity).json
                self.assertEqual(value["status"],"reporting");self.assertFalse(value["report_ready"])
                self.assertEqual(self.client.get("/api/jobs/"+identity+"/download").status_code,400)
            self.assertEqual(self.client.post("/api/analyze",json={"upload_id":self.upload_id,"schemas":["DM_CAD"]},headers=self.headers).status_code,409)
            gate.set()
            self.assertEqual(self.finish(identity)["status"],"complete")

    def test_cancel_in_second_group_keeps_first_group_and_partial_rows(self):
        self.gate=threading.Event();self.cancel_group=2
        identity=self.start(group_size=4)
        self.assertTrue(self.ready.wait(5))
        self.assertEqual(self.client.post("/api/jobs/"+identity+"/cancel",json={},headers=self.headers).json["status"],"cancelling")
        self.gate.set()
        job=self.finish(identity)
        self.assertEqual(job["status"],"cancelled");self.assertEqual(job["result_count"],6)
        self.assertEqual(len(self.calls),2);self.assertFalse(job["report_ready"])
        self.assertEqual(self.client.get("/api/jobs/"+identity+"/download").status_code,400)
        report=self.client.post("/api/jobs/"+identity+"/report",json={},headers=self.headers).json
        self.assertEqual((report["status"],report["report_ready"]),("cancelled",True))
        self.assertEqual(self.downloaded(identity),200)
        self.assertEqual(self.client.get("/api/uploads/"+self.upload_id+"/rows").json["searched"],6)

    def test_since_slice_latest_and_other_owner(self):
        self.assertEqual(self.client.get("/api/jobs/latest").json,{})
        first=self.finish(self.start(row_ids=self.row_ids()[:3]))
        second=self.finish(self.start())
        sliced=self.client.get("/api/jobs/"+second["id"]+"?since=7").json
        self.assertEqual(sliced["result_count"],10);self.assertEqual(len(sliced["results"]),3)
        self.assertEqual(sliced["results"][0]["row_id"],second["results"][7]["row_id"])
        self.assertEqual(sliced["upload_id"],self.upload_id)
        self.assertEqual(self.client.get("/api/jobs/latest").json,{"id":second["id"],"status":"complete","upload_id":self.upload_id})
        self.assertNotEqual(first["id"],second["id"])
        other=self.app.test_client();other.get("/")
        self.assertEqual(other.get("/api/jobs/"+second["id"]).status_code,404)
        self.assertEqual(other.get("/api/uploads/"+self.upload_id+"/rows").status_code,404)
        self.assertEqual(other.get("/api/jobs/latest").json,{})
        for private in ("cancel_event","report_lock","partial_results","owner","CREDENTIAL_SENTINEL"):
            self.assertNotIn(private,self.client.get("/api/jobs/"+second["id"]).text)

    def test_finished_jobs_beyond_limit_evict_oldest(self):
        ids=self.row_ids();identities=[]
        for n in range(21):identities.append(self.finish(self.start(row_ids=[ids[n%10]]))["id"])
        jobs=self.app.extensions["riskmatch"]["jobs"]
        self.assertEqual(len(jobs),20);self.assertNotIn(identities[0],jobs)
        self.assertFalse((self.root/identities[0]).exists())

    def test_index_endpoints_use_owner_resources(self):
        self.assertEqual(self.client.get("/api/index/status").json,{"state":"idle"})
        off=self.client.post("/api/index/start",json={"schemas":["DM_CAD"],"settings":{"use_models":False}},headers=self.headers)
        self.assertEqual(off.status_code,400)
        started=self.client.post("/api/index/start",json={"schemas":["DM_CAD"],"settings":EMBEDDING},headers=self.headers)
        self.assertEqual(started.status_code,202,started.json)
        self.assertEqual(started.json["state"],"running");self.assertEqual(started.json["index_coverage"],0.5)
        self.assertEqual(self.client.get("/api/index/status").json["state"],"running")
        self.assertEqual(self.client.post("/api/index/stop",json={},headers=self.headers).json["state"],"stopped")
        refresh=self.client.post("/api/catalog/refresh",json={"schemas":["DM_CAD"]},headers=self.headers)
        self.assertEqual(refresh.json,{"columns":2,"source":"oracle"})
        # v0.5 key is pipeline.catalog_key(oracle, credentials): an opaque digest shared by start and refresh.
        catalog_calls=[c for c in self.resources[0].calls if c[0]=="catalog"]
        self.assertEqual([c[2] for c in catalog_calls],[False,True])
        self.assertEqual(len({c[1] for c in catalog_calls}),1)
        self.assertNotIn("EMBED_SECRET",catalog_calls[0][1])
        self.assertNotIn("EMBED_SECRET",started.text)
        self.client.post("/api/disconnect",json={},headers=self.headers)
        self.assertTrue(self.resources[0].closed)
        self.assertEqual(self.client.get("/api/index/status").json,{"state":"idle"})

    def test_worker_requires_full_index_only_with_embedding_models(self):
        self.finish(self.start(row_ids=self.row_ids()[:1],settings={"use_models":False}))
        self.finish(self.start(row_ids=self.row_ids()[:1],settings=EMBEDDING))
        self.assertEqual([c["require_full_index"] for c in self.calls],[False,True])

    def test_index_status_complete_only_when_done_and_full(self):
        started=self.client.post("/api/index/start",json={"settings":EMBEDDING},headers=self.headers)
        self.assertFalse(started.json["complete"])
        indexer=self.resources[0].last
        indexer.state="done"
        self.assertFalse(self.client.get("/api/index/status").json["complete"])
        indexer.index.coverage=1.0
        self.assertTrue(self.client.get("/api/index/status").json["complete"])
        indexer.index.tables_missing=1
        self.assertFalse(self.client.get("/api/index/status").json["complete"])

    def test_index_start_does_not_stop_same_cached_indexer(self):
        self.client.post("/api/index/start",json={"settings":EMBEDDING},headers=self.headers)
        indexer=self.resources[0].last
        self.resources[0].indexer=lambda snapshot,ctx,settings:indexer
        again=self.client.post("/api/index/start",json={"settings":EMBEDDING},headers=self.headers)
        self.assertEqual(again.status_code,202,again.json)
        self.assertEqual((indexer.stops,indexer.state),(0,"running"))

    def test_job_waits_for_running_index_without_holding_executor(self):
        self.client.post("/api/index/start",json={"settings":EMBEDDING},headers=self.headers)
        indexer=self.resources[0].last
        with patch("riskmatch.web.INDEX_POLL_SECONDS",0.01):
            identity=self.start(row_ids=self.row_ids()[:1],settings=EMBEDDING)
            deadline=time.monotonic()+5
            while self.client.get("/api/jobs/"+identity).json["progress"].get("phase")!="index_wait":
                self.assertLess(time.monotonic(),deadline);time.sleep(.01)
            # Tek analiz yuvası boş kalır; indeks bitmeden pipeline çağrılmaz.
            self.assertEqual(self.app.extensions["riskmatch"]["executor"].submit(lambda:1).result(timeout=2),1)
            self.assertEqual(self.calls,[])
            self.assertIn("İndeks tamamlanması bekleniyor",self.client.get("/api/jobs/"+identity).json["progress"]["stage"])
            indexer.state="done"
            self.assertEqual(self.finish(identity)["status"],"complete")
        self.assertEqual([c["require_full_index"] for c in self.calls],[True])

    def test_catalog_refresh_stops_and_forgets_indexer(self):
        self.client.post("/api/index/start",json={"settings":EMBEDDING},headers=self.headers)
        indexer=self.resources[0].last
        refresh=self.client.post("/api/catalog/refresh",json={"schemas":["NOT_VISIBLE"]},headers=self.headers)
        self.assertEqual(refresh.status_code,200,refresh.json)
        self.assertEqual(indexer.state,"stopped")
        self.assertEqual(self.client.get("/api/index/status").json,{"state":"idle"})

    def test_client_schemas_are_ignored_for_index_and_connection_keeps_fixed_scope(self):
        started=self.client.post("/api/index/start",json={"schemas":["NOT_VISIBLE"],"settings":EMBEDDING},headers=self.headers)
        self.assertEqual(started.status_code,202,started.json)
        status=self.client.get("/api/connection").json
        self.assertEqual((status["scan_schemas"],status["missing_schemas"]),(["DM_CAD"],["DM","EDW"]))
        self.assertIn("karar verilemedi",status["warning"])

    def pause_first_group(self,number):
        if number==1:self.ready.set();self.gate.wait(5)

    def info_sheet(self,identity):
        wb=load_workbook(self.root/identity/"sonuclar.xlsx",read_only=True)
        try:return {r[0]:r[1] for r in wb["Çalışma bilgisi"].iter_rows(min_row=2,values_only=True)}
        finally:wb.close()

    def test_second_analyze_while_running_keeps_workspace_and_job_mapping(self):
        self.gate=threading.Event();self.before_group=self.pause_first_group
        space=self.app.extensions["riskmatch"]["uploads"][self.upload_id]
        identity=self.start(group_size=5)
        self.assertTrue(self.ready.wait(5))
        original=dict(space["mapping"]);results=dict(space["results"])
        other={"name":"Açıklama","description":"Alan Adı"}
        response=self.client.post("/api/analyze",json={"upload_id":self.upload_id,"schemas":["DM_CAD"],"mapping":other},headers=self.headers)
        self.assertEqual(response.status_code,409)
        self.assertEqual((space["mapping"],space["results"]),(original,results))
        space["mapping"]=other  # Later groups must still use the job's snapshot.
        self.gate.set()
        self.assertEqual(self.finish(identity)["status"],"complete")
        self.assertEqual([c["mapping"] for c in self.calls],[original,original])

    def test_cancelled_job_report_says_incomplete(self):
        self.gate=threading.Event();self.cancel_group=2
        complete={"Tamamlanma":"Tamamlandı","Hatalar":[],"Uyarılar":[]}
        self.infos={1:complete,2:complete}
        identity=self.start(row_ids=self.row_ids()[:8],group_size=4)
        self.assertTrue(self.ready.wait(5))
        self.client.post("/api/jobs/"+identity+"/cancel",json={},headers=self.headers);self.gate.set()
        self.assertEqual(self.finish(identity)["status"],"cancelled")
        report=self.client.post("/api/jobs/"+identity+"/report",json={},headers=self.headers).json
        self.assertTrue(report["report_ready"])
        self.assertEqual(self.info_sheet(identity)["Tamamlanma"],"Eksik")

    def test_job_run_info_merges_groups(self):
        self.infos={1:{"Tamamlanma":"Eksik / teknik hata","Hatalar":["H1"],"Uyarılar":["U1"]},
                    2:{"Tamamlanma":"Tamamlandı","Hatalar":[],"Uyarılar":[]}}
        job=self.finish(self.start(group_size=5))
        self.assertEqual(job["status"],"complete")
        self.assertEqual(job["run_info"]["Tamamlanma"],"Eksik")
        self.assertEqual((job["run_info"]["Uyarılar"],job["run_info"]["Hatalar"]),(["U1"],["H1"]))
        self.assertEqual(self.info_sheet(job["id"])["Tamamlanma"],"Eksik")

    def test_model_change_in_last_group_marks_every_row(self):
        def info(check):
            return {"Tamamlanma":"Tamamlandı" if check=="same" else "Eksik / teknik hata","Hatalar":[],"Uyarılar":[],
                    "Model doğrulaması":{"embedding":{"end_of_job_check":check}}}
        self.infos={1:info("same"),2:info("same"),3:info("changed")}
        job=self.finish(self.start(group_size=4))
        self.assertEqual(job["result_count"],10);self.assertEqual(job["run_info"]["Tamamlanma"],"Eksik")
        self.assertTrue(all(r["model_changed"] and r["reasons"].count(MODEL_CHANGED_REASON)==1 for r in job["results"]))
        space=self.app.extensions["riskmatch"]["uploads"][self.upload_id]
        self.assertTrue(all(r["model_changed"] for r in space["results"].values()))
        self.assertTrue(all(r["model_changed"] for r in space["raw"].values()))
        self.assertEqual(job["run_info"]["Model doğrulaması"]["embedding"]["end_of_job_check"],"changed")
        self.assertEqual(job["run_info"]["Sonuç dağılımı"],{"found":10})

    def test_disconnect_during_job_defers_resources_close(self):
        self.gate=threading.Event();self.before_group=self.pause_first_group
        identity=self.start(group_size=5)
        self.assertTrue(self.ready.wait(5))
        self.assertEqual(self.client.post("/api/disconnect",json={},headers=self.headers).status_code,200)
        self.assertFalse(self.resources[0].closed)
        self.gate.set()
        self.assertEqual(self.finish(identity)["status"],"complete")
        self.assertEqual(len(self.resources),1)
        self.assertTrue(all(c["resources"] is self.resources[0] for c in self.calls))
        self.assertEqual(self.resources[0].closes,[True])

    def test_failure_after_first_group_keeps_rows_and_allows_report(self):
        def fail(number):
            if number==2:raise OSError("disk SECRET_PATH")
        self.before_group=fail
        job=self.finish(self.start(group_size=5))
        self.assertEqual(job["status"],"failed");self.assertEqual(job["result_count"],5)
        self.assertNotIn("SECRET_PATH",job["error"])
        report=self.client.post("/api/jobs/"+job["id"]+"/report",json={},headers=self.headers).json
        self.assertEqual((report["status"],report["report_ready"]),("failed",True))
        self.assertEqual(self.downloaded(job["id"]),200)
        self.assertEqual(self.info_sheet(job["id"])["Tamamlanma"],"Eksik")

    def test_report_retry_after_mapping_change_keeps_own_rows(self):
        ids=self.row_ids();written=[];real=excelio.write_results
        with patch("riskmatch.web.write_results",side_effect=ReportWriteError("locked")):
            first=self.finish(self.start(row_ids=ids))
        self.assertEqual(first["status"],"report_failed")
        other=copy.deepcopy(self.app.extensions["riskmatch"]["uploads"][self.upload_id]["mapping"])
        other["Değişkenler"]["columns"]["category"]=None
        self.finish(self.start(mapping=other,row_ids=ids[:1]))
        def capture(input_path,output_path,requests,results,run_info):
            written.append(len(results));return real(input_path,output_path,requests,results,run_info)
        with patch("riskmatch.web.write_results",side_effect=capture):
            retry=self.client.post("/api/jobs/"+first["id"]+"/report",json={},headers=self.headers).json
        self.assertTrue(retry["report_ready"]);self.assertEqual(written,[10])
        self.assertEqual(self.info_sheet(first["id"])["Aranan satır"],10)

    def test_index_start_after_disconnect_is_not_started(self):
        connections=self.app.extensions["riskmatch"]["connections"];made=[]
        def racing(resource,snapshot,ctx,settings):
            connections.clear();made.append(FakeIndexer());return made[0]
        with patch.object(FakeResources,"indexer",racing):
            response=self.client.post("/api/index/start",json={"schemas":["DM_CAD"],"settings":EMBEDDING},headers=self.headers)
        self.assertEqual(response.status_code,409);self.assertEqual(made[0].state,"idle")

    def test_sweep_removes_only_old_job_folders(self):
        old=time.time()-2*86400
        for name in ("catalog","resources","a"*32):
            (self.root/name).mkdir();os.utime(self.root/name,(old,old))
        (self.root/"vectors.sqlite").write_bytes(b"x")
        create_app(self.root,analyzer=self.analyze)
        self.assertTrue((self.root/"catalog").is_dir());self.assertTrue((self.root/"resources").is_dir())
        self.assertTrue((self.root/"vectors.sqlite").exists());self.assertFalse((self.root/("a"*32)).exists())


if __name__=="__main__":unittest.main()

"""A valid bulk workbook must not look analyzed when Oracle inventory fails."""
from copy import deepcopy
import io
import json
import time
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import test_web
from test_adapters import FakeConnection
from riskmatch.adapters.oracle import OracleAdapter


ROW_COUNT = 273
SCHEMAS = ["RISK_DWH", "RISK_SECOND", "RISK_THIRD"]
MODEL_SETTINGS = {
    "use_models": True,
    "search_mode": "full", "llm_review": True,
    "llm": {"endpoint": "https://llm.invalid/v1/chat/completions",
            "model": "qwen38-27b", "model_version": "test-v1", "contract": "openai_chat",
            "api_key": "LLM_SECRET_SENTINEL"},
    "embedding": {"endpoint": "https://embedding.invalid/v1/embeddings",
                  "model": "", "model_version": "", "model_selection": "server_default",
                  "dimensions": 2, "max_batch": 1, "request_template": {"input": "{texts}"},
                  "response_path": "data[*].embedding", "api_key": "EMBED_SECRET_SENTINEL"},
}


def bulk_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Degiskenler"
    sheet.append(["Kategori", "Alan Adı", "Alan tipi", "Açıklama", "Halihazırda bildiriliyor mu"])
    for number in range(1, ROW_COUNT + 1):
        sheet.append(["Risk", "HAV_LM_" + str(number), "NUMBER",
                      "Son ay havale kisi sayisi", "Evet"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream


class TimeoutCursor:
    def __init__(self):
        self.sql = ""
        self.position = 0

    def execute(self, sql, binds=None):
        self.sql = sql
        self.position = 0
        if "ALL_TAB_COLUMNS" in sql:
            raise RuntimeError("DPY-4024: PRIVATE_DRIVER_DETAIL password=SECRET_SENTINEL")

    def fetchone(self):
        return ("RISK", "PDB1")

    def fetchmany(self, size):
        rows = self.fetchall()
        page = rows[self.position:self.position + size]
        self.position += len(page)
        return page

    def fetchall(self):
        if "ALL_OBJECTS" in self.sql:
            return [("CUSTOMER_METRIC", "TABLE")]
        return []


class TimeoutConnection(FakeConnection):
    def __init__(self):
        self.cur = TimeoutCursor()
        self.closed = False


class BulkModelService:
    def __init__(self):
        self.embedding_calls = 0
        self.llm_calls = 0

    def open(self, request, timeout):
        payload = json.loads(request.data)
        if request.full_url == MODEL_SETTINGS["embedding"]["endpoint"]:
            if set(payload) != {"input"} or len(payload["input"]) != 1:
                raise AssertionError("The model-free single-input contract changed")
            self.embedding_calls += 1
            response = {"data": [{"embedding": [1.0, 0.0]}]}
        elif request.full_url == MODEL_SETTINGS["llm"]["endpoint"]:
            self.llm_calls += 1
            context = json.loads(payload["messages"][-1]["content"])
            candidate = context["candidates"][0]
            response = {"choices": [{"message": {"content": json.dumps({"assessments": [
                {"catalog_id": candidate["catalog_id"], "evidence": candidate["column"],
                 "note": "Source evidence checked"}]})}}]}
        else:
            raise AssertionError("Unexpected HTTP endpoint")
        return io.BytesIO(json.dumps(response).encode("utf-8"))


class CatalogFailureWebTests(unittest.TestCase):
    setUp = test_web.WebTests.setUp
    tearDown = test_web.WebTests.tearDown

    def connect(self):
        connection = FakeConnection()
        connection.cur.fetchmany = lambda count: [(schema,) for schema in SCHEMAS]
        with patch.object(OracleAdapter, "_connect", return_value=connection):
            result = self.client.post("/api/connect", json={"user": "test-user",
                                      "password": "ORACLE_SECRET_SENTINEL"}, headers=self.headers)
        self.assertEqual(result.status_code, 200, result.json)
        self.assertTrue(result.json["ok"], result.json)
        self.assertEqual(result.json["schemas"], SCHEMAS)

    def upload_bulk(self):
        result = self.client.post("/api/upload", data={"file": (bulk_workbook(), "dictionary-273.xlsx")},
                                  headers=self.headers)
        self.assertEqual(result.status_code, 200, result.json)
        sheet = result.json["sheets"][0]
        self.assertEqual(sheet["rows"], ROW_COUNT)
        self.assertEqual(sheet["suggested"]["name"], 2)
        self.assertEqual(sheet["suggested"]["description"], 4)
        return result.json["id"], {sheet["name"]: {"header_row": sheet["header_row"],
                                                   "columns": sheet["suggested"]}}

    def finish(self, identity):
        deadline = time.monotonic() + 60
        while time.monotonic() < deadline:
            result = self.client.get("/api/jobs/" + identity)
            if result.json["status"] not in ("running", "cancelling", "reporting"):
                return result.json
            time.sleep(.02)
        self.fail("Bulk workbook job did not finish")

    def test_catalog_timeout_keeps_valid_workbook_at_zero_analyzed_rows(self):
        identity, mapping = self.upload_bulk()
        connections = []

        def connect():
            connection = TimeoutConnection()
            connections.append(connection)
            return connection

        with patch.object(OracleAdapter, "_connect", side_effect=connect), \
                patch("urllib.request.build_opener", side_effect=AssertionError("No model HTTP call is allowed")) as model_http:
            response = self.client.post("/api/analyze", json={
                "schemas": SCHEMAS, "upload_id": identity, "mapping": mapping,
                "settings": deepcopy(MODEL_SETTINGS)}, headers=self.headers)
            self.assertEqual(response.status_code, 202, response.json)
            result = self.finish(response.json["id"])

        self.assertEqual(result["status"], "failed", result.get("error", result.get("progress")))
        self.assertEqual(result["failure_stage"], "catalog")
        self.assertEqual(result["code"], "DPY-4024")
        self.assertEqual(result["progress"]["completed"], 0)
        self.assertEqual(result["progress"]["total"], ROW_COUNT)
        self.assertFalse(result.get("results"))
        self.assertIn("DPY-4024", result["error"])
        self.assertIn(str(ROW_COUNT), result["error"])
        self.assertIn("Excel", result["error"])
        self.assertIn("katalog", result["error"].casefold())
        self.assertNotIn("SECRET_SENTINEL", json.dumps(result))
        self.assertNotIn("PRIVATE_DRIVER_DETAIL", json.dumps(result))
        model_http.assert_not_called()
        self.assertTrue(connections)
        self.assertTrue(all(connection.closed for connection in connections))
        self.assertEqual(self.client.get("/api/jobs/" + result["id"] + "/download").status_code, 400)
        self.assertFalse(list((self.root / result["id"]).rglob("sonuclar.xlsx")))

    def test_bulk_excel_runs_all_rows_through_both_real_adapters(self):
        identity, mapping = self.upload_bulk()
        service = BulkModelService()
        with patch.object(OracleAdapter, "_connect", return_value=FakeConnection()), \
                patch("urllib.request.build_opener", return_value=service):
            response = self.client.post("/api/analyze", json={
                "schemas": ["RISK_DWH"], "upload_id": identity, "mapping": mapping,
                "settings": deepcopy(MODEL_SETTINGS)}, headers=self.headers)
            self.assertEqual(response.status_code, 202, response.json)
            result = self.finish(response.json["id"])

        self.assertEqual(result["status"], "complete", result.get("error", result.get("progress")))
        self.assertEqual(result["run_info"]["Tamamlanma"], "Tamamlandı")
        self.assertEqual(result["progress"]["completed"], ROW_COUNT)
        self.assertEqual(result["progress"]["total"], ROW_COUNT)
        self.assertEqual(len(result["results"]), ROW_COUNT)
        # v0.5: LLM review only for rows with non-strong shown candidates; every requested review is made.
        self.assertEqual(service.llm_calls, sum(1 for row in result["results"] if row.get("llm_requested")))
        self.assertGreaterEqual(service.embedding_calls, ROW_COUNT + 1)
        for number, row in enumerate(result["results"], 1):
            self.assertEqual(row["errors"], [])
            self.assertEqual(row["request"]["name"], "HAV_LM_" + str(number))
            self.assertEqual(row["request"]["reported"], "Evet")
            if row.get("llm_requested"):
                self.assertEqual(row["model_review"]["model"], "qwen38-27b")
                self.assertTrue(row["model_review"]["assessments"])
        self.assertNotIn("SECRET_SENTINEL", json.dumps(result))
        response = self.client.get("/api/jobs/" + result["id"] + "/download")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data))
        self.assertEqual(workbook["Girdi 1 Degiskenler"].max_row, ROW_COUNT + 1)
        workbook.close()
        response.close()


if __name__ == "__main__":
    unittest.main()

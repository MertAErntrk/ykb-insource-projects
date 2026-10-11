"""HTTP workflow tests use injected adapters; no corporate connections or demo UI."""
from copy import deepcopy
import io
import json
from pathlib import Path
import tempfile
import time
import threading
import unittest
from unittest.mock import patch

from openpyxl import Workbook,load_workbook
from riskmatch.web import create_app
from riskmatch.web_services import clean_profile,inspect_workbook,validate_mapping,build_run,probe_model,normalize_dsn
from riskmatch.adapters._http import response_path
from riskmatch.errors import ConfigError,InputError,IntegrationError
from test_adapters import FakeConnection, ORACLE
from riskmatch.adapters.oracle import OracleAdapter

SETTINGS={"mode":"real","use_models":False,"oracle":{"dsn":"host.invalid:1521/service","user":"test-user","password":"SECRET_SENTINEL_93424","source_database":"risk","environment":"dev","container":"PDB1","schemas":["RISK_DWH"]}}

def workbook(headers=None):
    wb=Workbook();ws=wb.active;ws.title="Değişkenler"
    ws.append(headers or ["Kategori","Alan Adı","Alan tipi","Açıklama","Halihazırda bildiriliyor mu"])
    ws.append(["Risk","HAV_LM","NUMBER","Son ay havale kişi sayısı","Evet"])
    data=io.BytesIO();wb.save(data);wb.close();data.seek(0);return data


class WebTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.app=create_app(self.root);self.client=self.app.test_client();self.client.get("/")
        with self.client.session_transaction() as s:self.headers={"X-CSRF-Token":s["csrf"]}
    def tearDown(self):
        self.app.extensions["riskmatch"]["executor"].shutdown(wait=True)
        self.tmp.cleanup()
    def upload(self):
        result=self.client.post("/api/upload",data={"file":(workbook(),"dictionary.xlsx")},headers=self.headers)
        self.assertEqual(result.status_code,200,result.json)
        value=result.json;sheet=value["sheets"][0]
        return value["id"],{sheet["name"]:{"header_row":sheet["header_row"],"columns":sheet["suggested"]}}
    def finish(self,identity):
        deadline=time.monotonic()+15
        while time.monotonic()<deadline:
            result=self.client.get("/api/jobs/"+identity)
            if result.json["status"]!="running":return result.json
            time.sleep(.02)
        self.fail("Background job did not complete")
    def test_shell_real_only_and_security_headers(self):
        response=self.client.get("/")
        self.assertEqual(response.status_code,200)
        self.assertNotIn("demo",response.text.lower())
        self.assertIn("script-src 'self'",response.headers["Content-Security-Policy"])
        self.assertEqual(response.headers["Cache-Control"],"no-store")
        self.assertEqual(self.client.get("/",headers={"Host":"attacker.invalid"}).status_code,400)
    def test_csrf_and_cross_origin_rejected(self):
        self.assertEqual(self.client.post("/api/profile",json={}).status_code,403)
        self.assertEqual(self.client.post("/api/profile",json={},headers={**self.headers,"Origin":"https://attacker.invalid"}).status_code,403)
    def test_profile_omits_credentials_and_rejects_demo(self):
        data=deepcopy(SETTINGS);data["llm"]={"api_key":"secret","model":"id"}
        response=self.client.post("/api/profile",json=data,headers=self.headers)
        self.assertNotIn("password",response.text);self.assertNotIn("api_key",response.text)
        data["mode"]="demo"
        self.assertEqual(self.client.post("/api/profile",json=data,headers=self.headers).status_code,400)
        data={"embedding":{"request_template":{"Authorization":"secret"}}}
        with self.assertRaises(ConfigError):clean_profile(data)
    def test_five_headers_auto_detect_and_optional_fields(self):
        identity,mapping=self.upload()
        self.assertEqual(mapping["Değişkenler"]["columns"]["name"],2)
        requests=validate_mapping(self.app.extensions["riskmatch"]["uploads"][identity]["path"],mapping)
        self.assertEqual(requests[0]["name"],"HAV_LM")
        self.assertFalse(requests[0]["input_errors"])
    def test_custom_headers_mapped_by_position(self):
        path=self.root/"custom.xlsx";path.write_bytes(workbook(["A","B","C","D","E"]).getvalue())
        rows=validate_mapping(path,{"Değişkenler":{"header_row":1,"columns":{"name":2,"description":4}}})
        self.assertEqual(rows[0]["name"],"HAV_LM")
        self.assertEqual(rows[0]["original"]["E"],"Evet")
    def test_mapping_rejects_duplicate_or_missing_columns(self):
        path=self.root/"input.xlsx";path.write_bytes(workbook().getvalue())
        for columns in ({"name":2},{"name":2,"description":2},{"name":True,"description":4}):
            with self.assertRaises(InputError):validate_mapping(path,{"Değişkenler":{"header_row":1,"columns":columns}})
    def test_invalid_file_and_oversize_rejected(self):
        for name,data in (("input.xls",b"x"),("input.xlsx",b"not zip")):
            response=self.client.post("/api/upload",data={"file":(io.BytesIO(data),name)},headers=self.headers)
            self.assertEqual(response.status_code,400)
        self.app.config["MAX_CONTENT_LENGTH"]=10
        self.assertEqual(self.client.post("/api/upload",data=b"x"*100,headers=self.headers).status_code,413)
    def test_upload_ownership_and_path_traversal(self):
        identity,mapping=self.upload();other=self.app.test_client();other.get("/")
        with other.session_transaction() as s:headers={"X-CSRF-Token":s["csrf"]}
        self.assertEqual(other.post("/api/analyze",json={"upload_id":identity,"mapping":mapping,"settings":SETTINGS},headers=headers).status_code,404)
        self.assertEqual(self.client.get("/api/jobs/unknown/download").status_code,404)
    def test_live_pipeline_returns_source_and_download_without_secrets(self):
        identity,mapping=self.upload()
        connection=FakeConnection()
        with patch("riskmatch.adapters.oracle.OracleAdapter._connect",return_value=connection),patch("riskmatch.adapters.qwen.post_json",side_effect=AssertionError("Models disabled")):
            response=self.client.post("/api/analyze",json={"upload_id":identity,"mapping":mapping,"settings":SETTINGS},headers=self.headers)
            self.assertEqual(response.status_code,202,response.json)
            result=self.finish(response.json["id"])
        self.assertEqual(result["status"],"complete")
        self.assertEqual(result["results"][0]["physical_name"],"present")
        candidate=result["results"][0]["candidates"][0]
        self.assertEqual(candidate["source"]["object"],"CUSTOMER_METRIC")
        self.assertEqual(candidate["source"]["column"],"HAV_LM")
        self.assertNotEqual(result["results"][0]["semantic_status"],"supported")
        self.assertEqual(result["run_info"]["Arama yaklaşımı"],"bm25")
        response=self.client.get("/api/jobs/"+result["id"]+"/download")
        self.assertEqual(response.status_code,200)
        output=load_workbook(io.BytesIO(response.data));heads=[c.value for c in output["Adaylar"][1]]
        self.assertIn("Kolon açıklaması",heads)
        self.assertEqual(output["Girdi 1 Değişkenler"].cell(2,5).value,"Evet");output.close();response.close()
        for path in self.root.rglob("*.json"):
            self.assertNotIn(SETTINGS["oracle"]["password"],path.read_text("utf-8"))
        self.assertNotIn(SETTINGS["oracle"]["password"],json.dumps(result))
        other=self.app.test_client();other.get("/")
        self.assertEqual(other.get("/api/jobs/"+result["id"]).status_code,404)
    def test_oracle_failure_stays_incomplete(self):
        identity,mapping=self.upload()
        with patch("riskmatch.adapters.oracle.OracleAdapter._connect",side_effect=IntegrationError("SECRET_RAW_DRIVER")):
            response=self.client.post("/api/analyze",json={"upload_id":identity,"mapping":mapping,"settings":SETTINGS},headers=self.headers)
            result=self.finish(response.json["id"])
        self.assertFalse(result["run_info"]["Tarama tamamlandı"])
        self.assertEqual(result["results"][0]["physical_name"],"unknown")
        self.assertNotIn("SECRET_RAW_DRIVER",json.dumps(result))
    def test_oracle_probe_closes_and_finds_scope(self):
        connection=FakeConnection();connection.cur.fetchmany=lambda count:[("RISK_DWH",)]
        with patch("riskmatch.adapters.oracle.OracleAdapter._connect",return_value=connection):
            response=self.client.post("/api/probe/oracle",json=SETTINGS,headers=self.headers)
        self.assertEqual(response.json["container"],"PDB1")
        self.assertEqual(response.json["schemas"],["RISK_DWH"])
        self.assertTrue(connection.closed)
    def test_embedding_probe_reads_dimension_and_root_response(self):
        data={"embedding":{"endpoint":"https://embedding.invalid/embed","model":"id","model_version":"v1","request_template":{"inputs":"{texts}"},"response_path":"$"}}
        with patch("riskmatch.web_services.post_json",return_value=[[.1,.2,.3]]):
            self.assertEqual(probe_model(data,"embedding")["dimensions"],3)
        self.assertEqual(response_path({"data":[{"embedding":[1,2]}]},"data[*].embedding"),[[1,2]])
    def test_models_require_explicit_contract(self):
        data=deepcopy(SETTINGS);data["use_models"]=True
        with self.assertRaises(ConfigError):build_run(data)
    def test_dsn_alias_is_not_passed_to_driver(self):
        self.assertEqual(normalize_dsn("EDW = (DESCRIPTION=(ADDRESS=(HOST=host.invalid)))"),"(DESCRIPTION=(ADDRESS=(HOST=host.invalid)))")

    def test_catalog_limit_prevents_absence_claim(self):
        connection=FakeConnection()
        connection.cur.fetchmany=lambda count:connection.cur.fetchall()*2
        adapter=OracleAdapter({**ORACLE,"max_records":1},credentials={"user":"u","password":"p","dsn":"d"},connector=lambda **kwargs:connection)
        snapshot=adapter.scan()
        self.assertFalse(snapshot["complete"])
        self.assertEqual(len(snapshot["records"]),1)
        self.assertTrue(any("sınır" in error for error in snapshot["errors"]))

    def test_credentials_are_per_connection_not_global_environment(self):
        received=[]
        def connector(**kwargs):received.append(kwargs);return FakeConnection()
        with patch.dict("os.environ",{"TEST_ORACLE_USER":"env-user","TEST_ORACLE_PASSWORD":"env-pass","TEST_ORACLE_DSN":"env-dsn"}):
            for user in ("one","two"):
                adapter=OracleAdapter(ORACLE,connector=connector,credentials={"user":user,"password":user+"-secret","dsn":"host.invalid/service"})
                adapter._connect().close()
            import os
            self.assertEqual(os.environ["TEST_ORACLE_USER"],"env-user")
        self.assertEqual([row["user"] for row in received],["one","two"])

    def test_second_analysis_rejected_while_worker_active(self):
        identity,mapping=self.upload();gate=threading.Event();started=threading.Event()
        def delayed_scan(self):started.set();gate.wait(5);raise IntegrationError("test")
        with patch.object(OracleAdapter,"scan",delayed_scan):
            request={"upload_id":identity,"mapping":mapping,"settings":SETTINGS}
            first=self.client.post("/api/analyze",json=request,headers=self.headers)
            self.assertTrue(started.wait(2))
            try:
                second=self.client.post("/api/analyze",json=request,headers=self.headers)
                self.assertEqual(second.status_code,409)
            finally:gate.set()
            self.finish(first.json["id"])

    def test_model_enabled_web_pipeline_uses_both_explicit_services(self):
        identity,mapping=self.upload();data=deepcopy(SETTINGS);data["use_models"]=True
        data["llm"]={"endpoint":"https://llm.invalid/chat","model":"llm-id","model_version":"v1","contract":"openai_chat","api_key":"LLM_SECRET_SENTINEL"}
        data["embedding"]={"endpoint":"https://embed.invalid/embed","model":"embed-id","model_version":"v1","request_template":{"inputs":"{texts}"},"response_path":"$","dimensions":2,"api_key":"EMBED_SECRET_SENTINEL"}
        seen=[]
        def transport(config,payload,**kwargs):
            seen.append((config["endpoint"],kwargs["api_key"]))
            if "inputs" in payload:return [[.3,.7] for _ in payload["inputs"]]
            return {"choices":[{"message":{"content":json.dumps({"assessments":[]})}}]}
        with patch.object(OracleAdapter,"_connect",return_value=FakeConnection()),patch("riskmatch.adapters.qwen.post_json",side_effect=transport):
            response=self.client.post("/api/analyze",json={"upload_id":identity,"mapping":mapping,"settings":data},headers=self.headers)
            self.assertEqual(response.status_code,202,response.json)
            result=self.finish(response.json["id"])
        self.assertEqual(result["status"],"complete")
        self.assertIn(("https://llm.invalid/chat","LLM_SECRET_SENTINEL"),seen)
        self.assertIn(("https://embed.invalid/embed","EMBED_SECRET_SENTINEL"),seen)
        self.assertEqual(result["run_info"]["Arama yaklaşımı"],"hybrid")
        for path in self.root.rglob("*.json"):
            self.assertNotIn("SECRET_SENTINEL",path.read_text("utf-8"))


if __name__=="__main__":unittest.main()

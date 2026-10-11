"""Bound work before the first result, using real Excel and HTTP adapter contracts.

The network and Oracle inventory are controlled fixtures. This is not a claim
about a bank database's scan time or the quality of its live model responses.
"""

from copy import deepcopy
import io
import json
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch
from urllib.error import URLError

from openpyxl import Workbook, load_workbook

from riskmatch.errors import AnalysisCancelled
from riskmatch.pipeline import analyze
from riskmatch.semantics import assess_production, build_dependencies, parse_definition, prepare_production_catalog
import test_semantics


SCOPE = {"source_database": "test-db", "environment": "test", "container": "PDB1",
         "schemas": ["DM_CAD"]}
EMBEDDING_ENDPOINT = "https://embedding.internal.example/v1/embeddings"
LLM_ENDPOINT = "https://llm.internal.example/v1/chat/completions"
CATALOG_MARKER = "CATALOG_SCOPE_MARKER"


def fast_config(budget=256):
    return {
        "mode": "real", "allow_real_connections": True, "use_models": True,
        "llm_review": False,
        "retrieval": {"limit": 8, "expanded_limit": 24, "search_mode": "fast",
                      "candidate_limit": 24, "embedding_budget": budget},
        "oracle": {**SCOPE, "user_env": "TEST_USER", "password_env": "TEST_PASSWORD",
                   "dsn_env": "TEST_DSN"},
        "embedding": {"endpoint": EMBEDDING_ENDPOINT,
                      "allowed_hosts": ["embedding.internal.example"],
                      "model_selection": "server_default", "model": "", "model_version": "",
                      "dimensions": 4096, "request_template": {"input": "{texts}"},
                      "response_path": "data[*].embedding", "max_retries": 0},
        "llm": {"endpoint": LLM_ENDPOINT, "allowed_hosts": ["llm.internal.example"],
                "model": "test-llm", "model_version": "test-v1", "contract": "openai_chat",
                "max_retries": 0},
    }


def catalog(record_count=10000):
    records = []
    for number in range(record_count):
        records.append({**{key: SCOPE[key] for key in ("source_database", "environment", "container")},
                        "catalog_id": "column-{:05d}".format(number), "schema": "DM_CAD",
                        "object": "FEATURES_{:04d}".format(number // 20),
                        "column": "FIELD{:05d}".format(number), "data_type": "NUMBER",
                        # v0.5 document text is column + comment; the marker identifies catalog embeddings.
                        "column_comment": "Risk metric value {:05d} {}".format(number, CATALOG_MARKER),
                        "table_comment": CATALOG_MARKER, "metadata_version": "test-v1"})
    return {"records": records, "scope": deepcopy(SCOPE), "version": "fixed-catalog-v1",
            "complete": True, "errors": [], "captured_at": "2026-10-11T00:00:00Z"}


def write_workbook(path, names):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Degiskenler"
    sheet.append(["Kategori", "Alan Adı", "Alan tipi", "Açıklama", "Halihazırda bildiriliyor mu"])
    for name in names:
        sheet.append(["Risk", name, "NUMBER", "Requested risk metric value", "Evet"])
    workbook.save(path)
    workbook.close()


class ControlledModelService:
    def __init__(self, fail=False):
        self.catalog_texts = []
        self.query_texts = []
        self.llm_calls = 0
        self.attempts = 0
        self.fail = fail
        self.response_bytes = json.dumps({"data": [{"embedding": [1.0] + [0.0] * 4095}]}).encode("utf-8")

    def open(self, request, timeout):
        self.attempts += 1
        if request.full_url == LLM_ENDPOINT:
            self.llm_calls += 1
            raise AssertionError("Fast search must not request an LLM review unless explicitly enabled")
        if request.full_url != EMBEDDING_ENDPOINT:
            raise AssertionError("Unexpected test network request")
        payload = json.loads(request.data)
        if set(payload) != {"input"} or not isinstance(payload["input"], list) or len(payload["input"]) != 1:
            raise AssertionError("The institution-verified no-model/single-input contract changed")
        text = payload["input"][0]
        (self.catalog_texts if CATALOG_MARKER in text else self.query_texts).append(text)
        if self.fail:
            raise URLError("PRIVATE_TRANSPORT_DETAIL password=SECRET_SENTINEL")
        return io.BytesIO(self.response_bytes)


class RecoveringModelService(ControlledModelService):
    def __init__(self, embedding_down=False, llm_down=False):
        super().__init__(fail=embedding_down)
        self.llm_down = llm_down

    def open(self, request, timeout):
        if request.full_url == LLM_ENDPOINT:
            self.attempts += 1
            self.llm_calls += 1
            if self.llm_down:
                raise URLError("controlled LLM outage")
            return io.BytesIO(json.dumps({"choices": [{"message": {
                "content": '{"assessments": []}'}}]}).encode("utf-8"))
        return super().open(request, timeout)


class LargeCatalogPipelineTests(unittest.TestCase):
    def run_analysis(self, source, output, snapshot, service, config=None, **kwargs):
        with patch("riskmatch.adapters.oracle.OracleAdapter.scan", return_value=snapshot), \
                patch("urllib.request.build_opener", return_value=service):
            return analyze(source, output, mode="real", config_data=config or fast_config(),
                           state_dir=output.parent / "runtime", **kwargs)

    def test_10000_columns_and_273_rows_stream_before_global_embedding_finishes(self):
        snapshot = catalog()
        # Put exact matches at the end of the catalog, beyond every short list cap.
        names = ["FIELD{:05d}".format(9999 - number) for number in range(273)]
        service = ControlledModelService()
        streamed = []
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source, output = folder / "dictionary-273.xlsx", folder / "results.xlsx"
            write_workbook(source, names)
            started = time.monotonic()

            def on_result(row):
                streamed.append({"row": deepcopy(row), "catalog_calls": len(service.catalog_texts),
                                 "elapsed": time.monotonic() - started})

            result = self.run_analysis(source, output, snapshot, service, result_callback=on_result)
            elapsed = time.monotonic() - started
            self.assertTrue(output.is_file())
            workbook = load_workbook(output, read_only=True)
            self.assertEqual(workbook["Girdi 1 Degiskenler"].max_row, 274)
            workbook.close()
            self.assertFalse(list((folder / "runtime").rglob("vectors-*.json")))

        self.assertEqual(len(result["results"]), 273)
        self.assertEqual(len(streamed), 273)
        self.assertEqual(len({entry["row"]["row_id"] for entry in streamed}), 273)
        # v0.5: before the first row only the background sample (<=128) and that row's short list (<=10)
        # are embedded; the full index is the background indexer's job, never the analysis'.
        self.assertGreater(streamed[0]["catalog_calls"], 0)
        self.assertLessEqual(streamed[0]["catalog_calls"], 128 + 10)
        self.assertLess(streamed[0]["catalog_calls"], len(service.catalog_texts))
        self.assertLessEqual(len(service.catalog_texts), 128 + 10 * 273)
        self.assertLess(len(service.catalog_texts), len(snapshot["records"]))
        self.assertLessEqual(len(service.query_texts), 273 + 8 + 2)  # Plus fingerprint and end-of-job canaries.
        self.assertEqual(len(service.catalog_texts), len(set(service.catalog_texts)))
        self.assertEqual(service.llm_calls, 0)
        for name, row in zip(names, result["results"]):
            self.assertEqual(row["errors"], [])
            self.assertEqual(row["physical_name"], "present")
            self.assertTrue(any(candidate["source"]["column"] == name for candidate in row["candidates"]))
            self.assertNotEqual(row["semantic_status"], "not_found_in_scope")
            self.assertNotIn("model_review", row)
        # Partial semantic coverage is disclosed on the row instead of an embedding-budget note.
        self.assertEqual(result["results"][-1]["retrieval_info"]["semantic_scope"], "partial")
        self.assertTrue(any("anlamsal" in reason.casefold() for reason in result["results"][-1]["reasons"]))
        print("Large catalog fixture: 10000 columns / 273 requests; first row {:.3f}s; total {:.3f}s; "
              "catalog embeddings {}; query embeddings {}; LLM calls {}".format(
                  streamed[0]["elapsed"], elapsed, len(service.catalog_texts), len(service.query_texts), service.llm_calls))

    def test_unknown_model_vectors_are_reused_only_under_the_same_fingerprint(self):
        # v0.5: an unknown (server_default) model's stored vectors are reused when the canary fingerprint
        # namespace matches; a service producing different vectors gets a new namespace and no reuse.
        snapshot = catalog(128)
        service = ControlledModelService()

        def namespace(result):
            return result["run_info"]["Model doğrulaması"]["embedding"]["fingerprint"]["namespace"]

        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source, output = folder / "dictionary.xlsx", folder / "results.xlsx"
            write_workbook(source, ["FIELD00127", "FIELD00126"])
            first = self.run_analysis(source, output, snapshot, service)
            first_texts = list(service.catalog_texts)
            service.catalog_texts.clear()
            second = self.run_analysis(source, output, snapshot, service, resume=True)
            self.assertTrue(first_texts)
            self.assertEqual(service.catalog_texts, [])
            self.assertEqual(namespace(second), namespace(first))
            changed = ControlledModelService()
            changed.response_bytes = json.dumps({"data": [{"embedding": [0.0, 1.0] + [0.0] * 4094}]}).encode("utf-8")
            third = self.run_analysis(source, output, snapshot, changed, resume=True)
            self.assertNotEqual(namespace(third), namespace(first))
            self.assertTrue(changed.catalog_texts)
            self.assertEqual(third["run_info"]["Devam edilen satır"], 0)
            self.assertFalse(list((folder / "runtime").rglob("vectors-*.json")))

    def test_unavailable_embedding_keeps_lexical_evidence_and_exposes_degradation(self):
        snapshot = catalog(128)
        service = ControlledModelService(fail=True)
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source, output = folder / "dictionary.xlsx", folder / "results.xlsx"
            write_workbook(source, ["FIELD00127", "FIELD00126", "FIELD00125"])
            result = self.run_analysis(source, output, snapshot, service)
        self.assertGreater(service.attempts, 0)
        self.assertLessEqual(service.attempts, 2)
        self.assertEqual(service.llm_calls, 0)
        self.assertEqual(result["run_info"]["Tamamlanma"], "Eksik / teknik hata")
        for row in result["results"]:
            self.assertEqual(row["errors"], [])
            self.assertEqual(row["physical_name"], "present")
            self.assertTrue(row["candidates"])
            self.assertTrue(any("embedding" in reason.casefold() for reason in row["reasons"]))
        serialized = json.dumps(result, ensure_ascii=False)
        self.assertNotIn("PRIVATE_TRANSPORT_DETAIL", serialized)
        self.assertNotIn("SECRET_SENTINEL", serialized)

    def test_cancellation_stops_before_the_next_row_or_model_call(self):
        snapshot = catalog(128)
        service = ControlledModelService()
        streamed = []
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source, output = folder / "dictionary.xlsx", folder / "results.xlsx"
            write_workbook(source, ["FIELD00127", "FIELD00126", "FIELD00125"])

            def on_result(row):
                streamed.append((deepcopy(row), service.attempts))

            with self.assertRaises(AnalysisCancelled):
                self.run_analysis(source, output, snapshot, service,
                                  result_callback=on_result, cancelled=lambda: bool(streamed))
            self.assertEqual(len(streamed), 1)
            self.assertEqual(service.attempts, streamed[0][1])
            self.assertFalse(output.exists())

    def test_fast_search_does_not_require_an_unused_llm_configuration(self):
        service = ControlledModelService()
        config = fast_config()
        del config["llm"]
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source, output = folder / "dictionary.xlsx", folder / "results.xlsx"
            write_workbook(source, ["FIELD00000"])
            result = self.run_analysis(source, output, catalog(1), service, config=config)
        self.assertEqual(len(result["results"]), 1)
        self.assertEqual(service.llm_calls, 0)
        self.assertNotIn("llm", result["run_info"]["Model doğrulaması"])

    def retry_known_model_run(self, service, llm_review, runs=2):
        config = fast_config()
        config["llm_review"] = llm_review
        config["embedding"].update(model_selection="explicit", model="pinned", model_version="v1", max_batch=1)
        with tempfile.TemporaryDirectory() as directory, \
                patch("riskmatch.pipeline._source_version", return_value="fixed-test-source"):
            folder = Path(directory)
            source, output = folder / "dictionary.xlsx", folder / "results.xlsx"
            write_workbook(source, ["FIELD00000"])
            first = self.run_analysis(source, output, catalog(1), service, config=config)
            first_calls = service.attempts
            service.fail, service.llm_down = False, False
            for _ in range(runs - 1):
                before_last = service.attempts
                last = self.run_analysis(source, output, catalog(1), service, config=config, resume=True)
        return first, last, first_calls, before_last

    def test_stable_model_resume_retries_embedding_after_transient_failure(self):
        service = RecoveringModelService(embedding_down=True)
        first, second, first_calls, _ = self.retry_known_model_run(service, llm_review=False)
        self.assertEqual(first["results"][0]["errors"], [])
        # v0.5: the failed fingerprint call drops the run to name/text search; the run is marked incomplete.
        self.assertEqual(first["run_info"]["Tamamlanma"], "Eksik / teknik hata")
        self.assertTrue(first["run_info"]["Uyarılar"])
        self.assertEqual(second["run_info"]["Devam edilen satır"], 0)
        self.assertGreater(service.attempts, first_calls)
        self.assertFalse(second["results"][0]["retrieval_info"]["embedding_skipped"])
        # v0.5: the retried run embeds the query again (frozen index is empty, so no semantic channel hit is required).
        self.assertTrue(second["results"][0]["retrieval_info"]["query_embedded"])

    def test_stable_model_resume_retries_llm_review_after_transient_failure(self):
        service = RecoveringModelService(llm_down=True)
        first, second, first_calls, _ = self.retry_known_model_run(service, llm_review=True)
        self.assertEqual(first["results"][0]["errors"], [])
        self.assertNotIn("model_review", first["results"][0])
        self.assertEqual(second["run_info"]["Devam edilen satır"], 0)
        self.assertGreater(service.attempts, first_calls)
        self.assertEqual(service.llm_calls, 2)
        self.assertIn("model_review", second["results"][0])

    def test_complete_stable_model_results_still_resume_without_extra_calls(self):
        service = RecoveringModelService()
        # v0.5: the 2nd run re-searches because index coverage grew (first-run vectors were stored); once
        # coverage is stable the row resumes and only the 8 fingerprint canaries are sent (new Resources per run).
        first, third, first_calls, before_third = self.retry_known_model_run(service, llm_review=True, runs=3)
        self.assertIn("model_review", first["results"][0])
        self.assertEqual(third["run_info"]["Devam edilen satır"], 1)
        self.assertEqual(service.attempts - before_third, 8)


class ProductionPreparationTests(unittest.TestCase):
    def test_untrusted_requested_transaction_never_iterates_the_catalog(self):
        class UnreadableCatalog:
            def __iter__(self):
                raise AssertionError("No catalog row can establish an unknown requested transaction")

        request = test_semantics.request(definition=test_semantics.verified(concept="metric"))
        result = assess_production(request, {}, build_dependencies([request]), UnreadableCatalog())
        self.assertEqual(result["status"], "insufficient")
        self.assertEqual(result["dependencies"], [])

    def test_catalog_definition_parsing_is_once_per_run_and_movement_evidence_is_preserved(self):
        records = catalog(40)["records"]
        source = test_semantics.record(test_semantics.verified(grain="transaction", transaction="eft"), cid="source")
        records.append(source)
        request = test_semantics.request(definition=test_semantics.verified(transaction="eft"))
        graph = build_dependencies([request])
        expected = assess_production(request, {}, graph, records)
        with patch("riskmatch.semantics.parse_definition", wraps=parse_definition) as parser:
            context = prepare_production_catalog(records)
            self.assertEqual(parser.call_count, 40)
            for _ in range(273):
                result = assess_production(request, {}, graph, records, catalog_context=context)
                self.assertEqual(result, expected)
            self.assertEqual(parser.call_count, 40)
        self.assertEqual(result["dependencies"], ["source"])
        self.assertEqual(result["status"], "source_only")

    def test_catalog_context_is_scoped_to_its_own_snapshot(self):
        request = test_semantics.request(definition=test_semantics.verified(transaction="eft"))
        graph = build_dependencies([request])
        trusted = [test_semantics.record(test_semantics.verified(grain="transaction", transaction="eft"))]
        untrusted = [test_semantics.record(test_semantics.verified(grain="customer", transaction="eft"))]
        first = prepare_production_catalog(trusted)
        second = prepare_production_catalog(untrusted)
        self.assertEqual(assess_production(request, {}, graph, trusted, catalog_context=first)["status"], "source_only")
        self.assertEqual(assess_production(request, {}, graph, untrusted, catalog_context=second)["status"], "insufficient")

    def test_ratio_trust_checks_match_the_unprepared_path(self):
        items, records, results, graph = test_semantics.ProductionTests().fixture()
        context = prepare_production_catalog(records)
        self.assertEqual(assess_production(items[0], results, graph, records),
                         assess_production(items[0], results, graph, records, catalog_context=context))


if __name__ == "__main__":
    unittest.main()

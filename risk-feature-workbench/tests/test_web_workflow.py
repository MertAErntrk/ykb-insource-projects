"""Regression tests for built-in connections, schema selection and batch input."""
import io
import json
import unittest
from unittest.mock import patch
from openpyxl import load_workbook

from riskmatch.adapters.oracle import OracleAdapter, read_context, oracle_error
from riskmatch.deployment import ORACLE_DSN
from riskmatch.errors import OracleOperationError
from riskmatch.web_services import probe_oracle
from test_adapters import FakeConnection
import test_web
from test_web import SETTINGS


class WorkflowTests(unittest.TestCase):
    setUp=test_web.WebTests.setUp
    tearDown=test_web.WebTests.tearDown
    connect=test_web.WebTests.connect
    upload=test_web.WebTests.upload
    finish=test_web.WebTests.finish
    def test_login_uses_built_in_dsn_and_never_returns_password(self):
        seen=[]
        def connect(adapter):
            seen.append(adapter.credentials.copy())
            connection=FakeConnection()
            connection.cur.fetchmany=lambda count:[("RISK_DWH",)]
            return connection
        with patch.object(OracleAdapter,"_connect",connect):
            response=self.client.post("/api/connect",json={
                "user":"work-user","password":"PASSWORD_SENTINEL",
                "dsn":"ignored.invalid/service"},headers=self.headers)
        self.assertEqual(response.status_code,200)
        self.assertEqual(seen[0]["dsn"],ORACLE_DSN)
        self.assertNotIn("PASSWORD_SENTINEL",response.text)
        with self.client.session_transaction() as cookie:
            self.assertNotIn("PASSWORD_SENTINEL",json.dumps(dict(cookie)))
        self.assertFalse(any(self.root.rglob("*.json")))

    def test_manual_batch_returns_both_requests_and_sources(self):
        variables=[{"name":"HAV_LM","description":"Son ay havale kişi sayısı"},
                   {"name":"HAV_OTHER","description":"Havale yapan kişi sayısı"}]
        with patch.object(OracleAdapter,"_connect",return_value=FakeConnection()):
            response=self.client.post("/api/analyze",json={"input_type":"manual",
                "variables":variables,"schemas":["RISK_DWH"]},headers=self.headers)
            self.assertEqual(response.status_code,202,response.json)
            result=self.finish(response.json["id"])
        self.assertEqual(len(result["results"]),2)
        self.assertEqual([r["request"]["name"] for r in result["results"]],["HAV_LM","HAV_OTHER"])
        self.assertEqual(result["results"][0]["candidates"][0]["source"]["object"],"CUSTOMER_METRIC")
        download=self.client.get("/api/jobs/"+result["id"]+"/download")
        wb=load_workbook(io.BytesIO(download.data))
        self.assertEqual(wb["Sonuçlar"].max_row,3)
        wb.close();download.close()

    def test_standard_excel_can_run_without_mapping_payload(self):
        identity,_=self.upload()
        with patch.object(OracleAdapter,"_connect",return_value=FakeConnection()):
            response=self.client.post("/api/analyze",json={"upload_id":identity,
                "schemas":["RISK_DWH"]},headers=self.headers)
            self.assertEqual(response.status_code,202,response.json)
            result=self.finish(response.json["id"])
        self.assertEqual(result["results"][0]["request"]["reported"],"Evet")

    def test_schema_selection_required_and_must_come_from_login(self):
        for schemas in ([],["SYS"],["NOT_VISIBLE"],"RISK_DWH",[1]):
            response=self.client.post("/api/analyze",json={"input_type":"manual",
                "variables":[{"name":"x","description":"d"}],"schemas":schemas},headers=self.headers)
            self.assertEqual(response.status_code,400)
        self.assertFalse(self.app.extensions["riskmatch"]["jobs"])

    def test_multiple_schemas_passed_to_scan_and_empty_scope_stays_incomplete(self):
        seen=[]
        def connect(adapter):
            seen.append(adapter.schemas)
            return FakeConnection()
        with patch.object(OracleAdapter,"_connect",connect):
            response=self.client.post("/api/analyze",json={"input_type":"manual",
                "variables":[{"name":"x","description":"d"}],
                "schemas":["RISK_DWH","ANOTHER_SCHEMA"]},headers=self.headers)
            self.assertEqual(response.status_code,202,response.json)
            result=self.finish(response.json["id"])
        self.assertEqual(seen,[["ANOTHER_SCHEMA","RISK_DWH"]])
        self.assertFalse(result["run_info"]["Tarama tamamlandı"])
        self.assertEqual(result["results"][0]["physical_name"],"unknown")

    def test_connection_is_isolated_from_other_browser_session(self):
        other=self.app.test_client();other.get("/")
        self.assertFalse(other.get("/api/connection").json["ok"])
        with other.session_transaction() as cookie:headers={"X-CSRF-Token":cookie["csrf"]}
        other.post("/api/disconnect",json={},headers=headers)
        self.assertTrue(self.client.get("/api/connection").json["ok"])

    def test_manual_batch_validates_every_row_and_limit(self):
        for variables in ([],[{}],[{"name":"x","description":""}],
                          [{"name":"x","description":"d"}]*101):
            response=self.client.post("/api/analyze",json={"input_type":"manual",
                "variables":variables,"schemas":["RISK_DWH"]},headers=self.headers)
            self.assertEqual(response.status_code,400)
        self.assertFalse(self.app.extensions["riskmatch"]["jobs"])

    def test_logout_and_expiry_remove_credentials(self):
        with self.client.session_transaction() as cookie:owner=cookie["owner"]
        self.app.extensions["riskmatch"]["connections"][owner]["touched"]-=3601
        self.assertFalse(self.client.get("/api/connection").json["ok"])
        self.assertNotIn(owner,self.app.extensions["riskmatch"]["connections"])
        self.connect()
        self.client.post("/api/disconnect",json={},headers=self.headers)
        self.assertFalse(self.client.get("/api/connection").json["ok"])
        response=self.client.post("/api/analyze",json={"input_type":"manual",
            "variables":[{"name":"x","description":"d"}],"schemas":["RISK_DWH"]},headers=self.headers)
        self.assertEqual(response.status_code,400)

    def test_safe_error_code_survives_failed_scan_without_absence_claim(self):
        with patch.object(OracleAdapter,"_connect",side_effect=oracle_error(
                RuntimeError("ORA-01017 password=RAW_SECRET"),"oturum açma")):
            response=self.client.post("/api/analyze",json={"input_type":"manual",
                "variables":[{"name":"x","description":"d"}],"schemas":["RISK_DWH"]},headers=self.headers)
            result=self.finish(response.json["id"])
        self.assertIn("ORA-01017",json.dumps(result))
        self.assertNotIn("RAW_SECRET",json.dumps(result))
        self.assertEqual(result["results"][0]["physical_name"],"unknown")


class OracleDiagnosticTests(unittest.TestCase):
    def test_legacy_context_falls_back_only_for_unsupported_parameter(self):
        connection=FakeConnection();cursor=connection.cur
        execute=cursor.execute
        def legacy(sql,binds=None):
            if "CON_NAME" in sql:raise RuntimeError("ORA-02003 invalid USERENV parameter")
            execute(sql,binds)
        cursor.execute=legacy;cursor.fetchone=lambda:("LEGACYDB",)
        self.assertEqual(read_context(cursor),("LEGACYDB","LEGACYDB"))

    def test_context_privilege_error_is_not_hidden_by_fallback(self):
        connection=FakeConnection()
        connection.cur.execute=lambda *_:(_ for _ in ()).throw(RuntimeError("ORA-01031 RAW_SECRET"))
        with patch.object(OracleAdapter,"_connect",return_value=connection):
            report=probe_oracle(SETTINGS)
        self.assertFalse(report["ok"])
        self.assertTrue(report["connected"])
        self.assertIn("ORA-01031",report["error"])
        self.assertIn("bağlamı",report["error"].casefold())
        self.assertNotIn("RAW_SECRET",json.dumps(report))
        self.assertTrue(connection.closed)

    def test_schema_phase_has_its_own_code(self):
        connection=FakeConnection()
        connection.cur.fetchmany=lambda *_:(_ for _ in ()).throw(RuntimeError("ORA-00942 RAW_SECRET"))
        with patch.object(OracleAdapter,"_connect",return_value=connection):
            report=probe_oracle(SETTINGS)
        self.assertFalse(report["ok"])
        self.assertIn("ORA-00942",report["error"])
        self.assertIn("listeleme",report["error"])
        self.assertNotIn("RAW_SECRET",json.dumps(report))
        self.assertTrue(connection.closed)


if __name__=="__main__":unittest.main()

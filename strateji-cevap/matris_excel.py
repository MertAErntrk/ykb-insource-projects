"""
matris_excel.py — yetki matrisini Excel'den DÜZENE göre (renk değil) okuyup v2 matris JSON'u üretir.

Kullanım: python matris_excel.py --excel <yol> [--sayfa ...] [--cikti ...] [--config config.json]
          [--dogrula nominal=.. agirlikli=.. segment=.. beklenen=..] [--ayrinti]

Ekrana yalnız teşhis basılır (segment/satır sayısı, sayıya çevrilemeyen hücre adresleri, doğrulama sonucu);
hiçbir değer ya da unvan basılmaz. openpyxl yalnız bu betikte kullanılır (motor stdlib kalır).
"""
import argparse
import datetime
import json
import os
import re
import sys
from typing import Dict, List, Optional, Tuple

KOK = os.path.dirname(os.path.abspath(__file__))

_TR = str.maketrans({"İ": "I", "I": "I", "ı": "I", "i": "I", "Ş": "S", "ş": "S", "Ğ": "G", "ğ": "G",
                     "Ü": "U", "ü": "U", "Ö": "O", "ö": "O", "Ç": "C", "ç": "C"})
_KOSUL_RE = re.compile(r"nominal\s*(>=|<=|<|>)\s*(.+)", re.IGNORECASE)


def segment_adi(ad: str) -> str:
    """Türkçe harfler ASCII, büyük harf, boşluk ve '/' → '_'."""
    s = str(ad).strip().translate(_TR).upper()
    return re.sub(r"[\s/]+", "_", s)


def _bos(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _sayi(v) -> Tuple[Optional[float], bool]:
    """(değer, çevrilemedi). '-' / boş / '>' ile başlayan metin → (None, False)."""
    if _bos(v) or isinstance(v, bool):
        return None, False
    if isinstance(v, (int, float)):
        return (int(v) if float(v).is_integer() else float(v)), False
    t = str(v).strip()
    if t == "-" or t.startswith(">"):
        return None, False
    t = re.sub(r"(?i)usd|tl|[.\s,]", "", t)
    try:
        return int(t), False
    except ValueError:
        return None, True


def _kosul(v) -> Optional[dict]:
    """'Nominal < X' / '> X' / '>= X' → {'alan','op','deger'}; çözülemezse None."""
    if _bos(v):
        return None
    m = _KOSUL_RE.search(str(v))
    if not m or m.group(1) == "<=":
        return None
    deger, _ = _sayi(m.group(2))
    if deger is None:
        return None
    return {"alan": "nominal", "op": "<" if m.group(1) == "<" else ">=", "deger": deger}


def _birlesik_harita(ws, sutun: str, ilk: int, son: int) -> Dict[int, object]:
    """Satır → birleşik alanın sol üst hücre değeri (yalnız birleşik alanlar)."""
    from openpyxl.utils import column_index_from_string
    ci = column_index_from_string(sutun)
    h = {}
    for r in ws.merged_cells.ranges:
        if r.min_col <= ci <= r.max_col:
            deger = ws.cell(row=r.min_row, column=r.min_col).value
            for satir in range(max(r.min_row, ilk), min(r.max_row, son) + 1):
                h[satir] = deger
    return h


def _oku_ic(yol, sayfa="ST X1 TR", baslik_satiri=3, segment_sutun="B", unvan_sutun="C", kosul_sutun="D",
            agirlikli_sutun="E", nominal_sutun="F", unvan_seviye=None) -> Tuple[dict, List[str]]:
    import openpyxl
    wb = openpyxl.load_workbook(yol, data_only=True)
    if sayfa not in wb.sheetnames:
        raise ValueError("sayfa bulunamadı: {}".format(sayfa))
    ws = wb[sayfa]
    ilk, son = baslik_satiri + 1, ws.max_row
    birlesik = _birlesik_harita(ws, segment_sutun, ilk, son)
    gecersiz = []   # çevrilemeyen hücre adresleri
    segmentler = {}  # type: Dict[str, List[dict]]
    sirali = {}      # segment → {unvan: seviye}
    hesaplanan = {}  # type: Dict[str, int]
    onceki_seg = None
    onceki_unvan = None
    for r in range(ilk, son + 1):
        hucre = lambda s: ws["{}{}".format(s, r)]
        ham_seg = birlesik[r] if r in birlesik else hucre(segment_sutun).value
        if not _bos(ham_seg):
            onceki_seg = segment_adi(ham_seg)
        unvan = hucre(unvan_sutun).value
        kosul_ham = hucre(kosul_sutun).value
        if _bos(unvan) and _bos(kosul_ham):
            continue
        if onceki_seg is None:
            continue
        seg = onceki_seg
        satirlar = segmentler.setdefault(seg, [])
        sira_map = sirali.setdefault(seg, {})
        if _bos(unvan):
            if onceki_unvan is None or not satirlar:
                continue
            unvan = onceki_unvan
        else:
            unvan = str(unvan).strip()
        if unvan not in sira_map:
            # blok içindeki farklı unvan sırası: Otomatik Onay varsa 0'dan, yoksa 1'den
            oto_ilk = next(iter(sira_map), unvan) == "Otomatik Onay"
            if unvan == "Otomatik Onay":
                sira_map[unvan] = 0
            elif unvan_seviye and unvan in unvan_seviye:
                sira_map[unvan] = int(unvan_seviye[unvan])
            else:
                sira_map[unvan] = len(sira_map) + (0 if oto_ilk else 1)
        seviye = sira_map[unvan]
        hesaplanan.setdefault(unvan, seviye)
        onceki_unvan = unvan
        n, hn = _sayi(hucre(nominal_sutun).value)
        a, ha = _sayi(hucre(agirlikli_sutun).value)
        if hn:
            gecersiz.append(hucre(nominal_sutun).coordinate)
        if ha:
            gecersiz.append(hucre(agirlikli_sutun).coordinate)
        satir = {"sira": len(satirlar) + 1, "unvan": unvan, "seviye": seviye,
                 "nominal_ust": n, "agirlikli_ust": a}
        k = _kosul(kosul_ham)
        if k:
            satir["kosul"] = k
            if k["op"] == "<":
                satir["tavan"] = True
        satirlar.append(satir)
    segmentler = {s: v for s, v in segmentler.items() if v}
    sonuc_seviye = dict(hesaplanan)
    if unvan_seviye:
        sonuc_seviye.update(unvan_seviye)
    matris = {
        "surum": "excel-" + datetime.date.today().isoformat(),
        "kaynak": "excel",
        "para_birimi": "USD",
        "agirlikli_alan": "WEIGHTEDLIMWCASHCOLL",
        "otomatik_onay_seviyesi": 0,
        "unvan_seviye": sonuc_seviye,
        "segmentler": segmentler,
    }
    return matris, gecersiz


def oku(yol, sayfa="ST X1 TR", baslik_satiri=3, segment_sutun="B", unvan_sutun="C", kosul_sutun="D",
        agirlikli_sutun="E", nominal_sutun="F", unvan_seviye=None) -> dict:
    return _oku_ic(yol, sayfa, baslik_satiri, segment_sutun, unvan_sutun, kosul_sutun,
                   agirlikli_sutun, nominal_sutun, unvan_seviye)[0]


def _config(yol: Optional[str]) -> dict:
    if not yol:
        yol = os.path.join(KOK, "config.json")
        if not os.path.exists(yol):
            return {}
    with open(yol, encoding="utf-8") as f:
        return json.load(f)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Excel yetki matrisinden v2 matris JSON'u üretir")
    ap.add_argument("--excel", required=True)
    ap.add_argument("--sayfa", default="ST X1 TR")
    ap.add_argument("--cikti")
    ap.add_argument("--config")
    ap.add_argument("--dogrula", nargs="+", metavar="ANAHTAR=DEGER")
    ap.add_argument("--ayrinti", action="store_true")
    args = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except (AttributeError, ValueError):
        pass

    cfg = _config(args.config)
    matris, gecersiz = _oku_ic(args.excel, sayfa=args.sayfa, unvan_seviye=cfg.get("unvan_seviye") or None)
    cikti = args.cikti or cfg.get("matris_dosya") or os.path.join(KOK, "veri", "yetki_matrisi.json")
    klasor = os.path.dirname(os.path.abspath(cikti))
    os.makedirs(klasor, exist_ok=True)
    with open(cikti, "w", encoding="utf-8") as f:
        json.dump(matris, f, ensure_ascii=False, indent=2)

    seg = matris["segmentler"]
    print("segment: {}, satır: {}, çevrilemeyen: {}".format(
        len(seg), sum(len(v) for v in seg.values()), ", ".join(gecersiz) if gecersiz else "yok"))
    if args.ayrinti:
        for ad, v in seg.items():
            print("{}: {} satır".format(ad, len(v)))

    if args.dogrula:
        import matris as m
        d = dict(x.split("=", 1) for x in args.dogrula if "=" in x)
        try:
            hesap, _ = m.seviye_bul(matris, float(d["nominal"]), float(d["agirlikli"]), d.get("segment"))
            beklenen = int(d["beklenen"])
        except (KeyError, ValueError, m.MatrisHatasi):
            print("doğrulama kalmadı: girdi ya da hesap hatası")
            return 1
        if hesap == beklenen:
            print("doğrulama geçti")
        else:
            print("doğrulama kalmadı: hesap {}, beklenen {}".format(hesap, beklenen))
            return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())

import datetime as dt

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.styles import Font, PatternFill  # noqa: E402

import excel_dok  # noqa: E402


@pytest.fixture
def dosya(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Smile"
    ws.append(["Smile dökümü 2025"])                       # başlık üstü satır
    ws.append(["Açılış Tarihi", "Müşteri No", "Açıklama", "Durum", None])
    for i in range(30):
        ws.append([dt.datetime(2025, 1 + i % 12, 3), 10000001 + i,
                   "Müşteri 1234567{} için H221 neden çalıştı, kayıt 98765432".format(i % 10),
                   "Kapalı" if i % 3 else "Açık"])
    m = wb.create_sheet("Matris")
    m["A1"] = "Nominal Limit"
    m["A1"].font = Font(bold=True)
    m["B1"] = 5000000
    m["B1"].fill = PatternFill(patternType="solid", fgColor="FFFFFF00")
    m["C1"] = 3
    m["C1"].fill = PatternFill(patternType="solid", fgColor="FF00FF00")
    m["D1"] = "=B1*2"
    m["A2"] = "Müşteri 87654321 notu"
    m.merge_cells("A3:C3")
    m["A3"] = "Birleşik başlık"
    yol = tmp_path / "ornek.xlsx"
    wb.save(yol)
    return str(yol)


def test_yapi_maskeli_ve_icerik_yazmaz(dosya):
    out = excel_dok.dokum_yapi(dosya)
    assert "sayfalar: Smile, Matris" in out
    assert "başlık satırı: 2; veri satırı: 30" in out
    assert "başlık üstü satırlar: Smile dökümü 2025" in out
    assert "kimlik numarası gibi" in out
    assert "10000001" not in out and "98765432" not in out and "1234567" not in out
    assert "serbest metin" in out and "H221 neden" not in out
    assert "dağılım: Kapalı (20); Açık (10)" in out
    assert "tarih: en eski 2025-01-03" in out and "aylık: 2025-01 3" in out
    assert "(sütun 5)" not in out                           # tamamen boş sütun atlanır


def test_yapi_ornek_bayragi_maskeler(dosya):
    out = excel_dok.dokum_yapi(dosya, ornek=1)
    assert "örnek (maskeli): Müşteri ###### için H221 neden çalıştı, kayıt ######" in out


def test_hucreler_renk_formul_birlesik(dosya):
    out = excel_dok.dokum_hucreler(dosya)
    assert "B1 | 5000000 | dolgu=FFFFFF00" in out            # tutar maskelenmez, sarı dolgu
    assert "C1 | 3 | dolgu=FF00FF00" in out
    assert "A1 | Nominal Limit | kalın" in out
    assert "formül==B1*2" in out
    assert "birleşik alanlar: A3:C3" in out
    assert "A2 | Müşteri ###### notu" in out                 # metindeki numara maskelenir


def test_hucreler_degersiz_tutar_gizler(dosya):
    out = excel_dok.dokum_hucreler(dosya, degersiz=True)
    assert "B1 | sayı | dolgu=FFFFFF00" in out and "C1 | sayı | dolgu=FF00FF00" in out
    assert "5000000" not in out
    assert "formül==B1*#" in out                              # sabit gizli, hücre adresi kalır
    assert "A1 | Nominal Limit | kalın" in out                # başlık görünür


def test_hucreler_limit_notu(dosya):
    out = excel_dok.dokum_hucreler(dosya, limit=2)
    assert "dolu hücre daha yazılmadı (--limit 2)" in out


def test_main_dosya_yazar(dosya, tmp_path, capsys):
    cikti = str(tmp_path / "d.txt")
    assert excel_dok.main([dosya, "--cikti", cikti]) == 0
    with open(cikti, encoding="utf-8") as f:
        assert "# Yapı dökümü: ornek.xlsx" in f.read()
    assert excel_dok.main([str(tmp_path / "yok.xlsx")]) == 2
